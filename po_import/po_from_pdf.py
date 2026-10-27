"""供应商 Rechnung/Lieferschein PDF → Odoo 采购单导入表（填进 Odoo 导出的 purchase.order 模板）。

用法：
    python3 -m po_import.po_from_pdf <Lieferschein.pdf> <purchase.order 模板.xlsx>
    python3 -m po_import.po_from_pdf L.pdf T.xlsx --exclude <PZN> <PZN>
    python3 -m po_import.po_from_pdf L.pdf T.xlsx --pick <PZN>=__export__.product_product_<…>
    python3 -m po_import.po_from_pdf L.pdf T.xlsx --po P12380   # 模板含多张 PO 时指定（不给就自动选空白 PO）
    python3 -m po_import.po_from_pdf L.pdf T.xlsx -o /tmp/x      # 输出目录（默认 output/YYYYMMDD）

同构 `so_import/so_from_pdf.py`，规矩照搬：**PZN 是唯一连接键，绝不按品名兜底**；
匹配不到 / 命中多个产品就整张表不出、退出码 2，剔除（`--exclude`）与多选一（`--pick`）由用户定。
**采购只用基础品**：带渠道 / 组合装尾缀的候选先自动排除（`is_variant`），见 ISSUES [PO 导入] B。

填值口径（2026-10-20 用户定，与该供应商历史 PO 的手工口径一致）：
    Order Lines/Product/ID   ← 匹配到的外部 ID
    Order Lines/Quantity     ← PDF `Menge`
    Order Lines/Unit Price   ← PDF `AEP`（净价，不乘税）
    Order Lines/Taxes[/ID]   ← Code 含 `M`（MwSt. ermäßigt）→ 7% Vorsteuer，否则 19% Vorsteuer
    Vendor Reference         ← PDF 单号（模板有此列才填）
列名列序照抄模板；第一行带模板第 2 行的 ID / Order Reference / Vendor，其余行头信息列留空。
表头界面名与技术名（`id` / `name` / `order_line/product_qty` …）都认；产品列若是 `order_line/product_id/.id`
就写数据库 ID。

PDF 版式**只支持一家**（明细行 `PZN Menge Packung DarForm 名称 Code AEP AVP Gesamt`，
另认 0 元附件行 `PZN Menge 名称 0,00 0,00 0,00`（如 BEIPACK，照样进 PO、单价 0、按 19%），
底部 `Anz. Pos.` / 两档税基 / `Warenwert gesamt` / `Endsumme`）。解析护栏，任一不符就拒绝出表：
    · 逐行 AEP × Menge = Gesamt
    · Σ Gesamt = 印的 Warenwert gesamt = Endsumme
    · 行数 / 总件数 = 印的 Anz. Pos.（如 `9/98`）
    · Code 含 M 的行 Gesamt 合计 = 印的 Warenwert ermäß. MwSt，其余 = volle MwSt
      ——税率归类有独立校验；品名末尾恰是单个字母被误读成 Code 也会在这里被拦。
防重：Odoo 里已有 Vendor Reference = 本单号的 PO 就停。
贷项：页面含 `Gutschrift` 或 `Aus RG/LF … vom …` 引用行就停（退货不走 PO 导入）。

版式 B（第二家批发商的 Lieferschein，表头 `Menge Einheit Artikelbezeichnung PZN Symbol Lief./Pos. AVP AEP Wert`）
走 `parse_text_b`，见 ISSUES [PO 导入] E：Symbol 含 `*` → 7%（此版式 `M` 是 phonetischer Artikel，不是减税），
AVP 可缺；护栏同构——逐行 AEP × Menge = Wert、`n Positionen m Stück`、`Gesamtsumme EUR 7% 19% 总` 全对上；
页眉单据类型不是 LIEFERSCHEIN / RECHNUNG 就停。

版式 C（批发商门户里的 Lieferschein 网页，浏览器 Ctrl+S 存成 .html 后当第一个参数传入）走 `parse_html_c`，
见 ISSUES [PO 导入] G：表头 `Menge Artikelbezeichnung PZN Ek.Preis Zeilenwert MwSt Kz`，单价 = Ek.Preis，
税看 MwSt 列（7% / 19%），`Kz` 不读；底部只有 Gesamt (Netto)，护栏是逐行算术 + 总额；同 PZN 多行照原样分行。

模板含多张 PO（Odoo 一次导出多张空 PO）且没给 `--po` 时自动选空白 PO（Odoo 里 draft/sent、无行、无 Vendor Reference），
按模板行序取；输出目录里已为别的 PDF 出过表的 PO 跳过，同一 PDF 重跑沿用原 PO（`choose_po`，ISSUES [PO 导入] F）。
"""
import argparse
import datetime as dt
import os
import re
import sys
from decimal import Decimal
from html.parser import HTMLParser

for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8")

from openpyxl import load_workbook

from common.xlsx import unique_path
from common.po import _po_base_sku
from so_import.so_from_pdf import StopError, apply_exclude, de_num, pzn8, read_pdf

HEAD_COLS = ("ID", "Order Reference", "Vendor")
COL_VREF = "Vendor Reference"
COL_PID = "Order Lines/Product/ID"
COL_QTY = "Order Lines/Quantity"
COL_PRICE = "Order Lines/Unit Price"
COL_TOTAL_QTY = "Order Lines/Total Quantity"
COL_TAX_ID = "Order Lines/Taxes/ID"
COL_TAX = "Order Lines/Taxes"

#: 角色 → 模板里认的列名：(界面名, 技术名)。Odoo 导出可选「界面名」或「技术名」两种表头，导入两种都认
#: （2026-10-20 用户第 28 份模板就是技术名）。`pdbid` 是产品的数据库 ID（`.id`），只有技术名形式。
ALIASES = {
    "id": ("ID", "id"),
    "ref": ("Order Reference", "name"),
    "vendor": ("Vendor", "partner_id"),
    "vref": (COL_VREF, "partner_ref"),
    "pid": (COL_PID, "order_line/product_id/id"),
    "pdbid": (None, "order_line/product_id/.id"),
    "qty": (COL_QTY, "order_line/product_qty"),
    "total_qty": (COL_TOTAL_QTY, "order_line/product_uom_qty"),
    "price": (COL_PRICE, "order_line/price_unit"),
    "tax_id": (COL_TAX_ID, "order_line/taxes_id/id"),
    "tax": (COL_TAX, "order_line/taxes_id"),
}

#: 税：(外部 ID, 名称)。德国 SKR03 科目表自带的两档进项税，2026-10-20 在 Odoo 查实（id 24 / 23）。
TAX_REDUCED = ("l10n_de_skr03.1_tax_vst_7_skr03", "7% Vorsteuer")
TAX_FULL = ("l10n_de_skr03.1_tax_vst_19_skr03", "19% Vorsteuer")

_NUM = r"\d{1,3}(?:\.\d{3})*,\d{2}"
#: 明细行：PZN Menge Packung DarForm 名称 Code(一个或多个 * E M S T V) AEP AVP Gesamt
ROW_RE = re.compile(rf"^(\d{{6,8}}) (\d+) (\S+) (\S+) (.+?) ((?:[*EMSTV] )*[*EMSTV]) ({_NUM}) ({_NUM}) ({_NUM})$")
#: 0 元附件行（如 `05855555 1 B E I P A C K 0,00 0,00 0,00`）：无 Packung/DarForm/Code，三个金额全为 0。
#: 用户定照样进 PO（数量照 PDF、单价 0），见 ISSUES [PO 导入] C。
ZERO_RE = re.compile(r"^(\d{6,8}) (\d+) (.+?) 0,00 0,00 0,00$")
#: 页眉第一行：单号 页码
DOC_RE = re.compile(r"^(\d{7}) (\d+)$")
#: 底部合计行：…Anz. Pos.(行数/件数) volle MwSt € ermäß. MwSt € gesamt €（€ 常被抽成乱码，按任意非空白吃掉）
FOOT_RE = re.compile(rf"(?<!\S)(\d+)/(\d+) ({_NUM}) \S+ ({_NUM}) \S+ ({_NUM})(?: \S+)?$")
END_RE = re.compile(rf"Endsumme \*+ ({_NUM})")
#: 贷项（Gutschrift）信号：页尾单据类型（如 `Gutschrift Menge + Wert`；正常发货是 `Normal-Auftrag`），
#: 或行下引用原单 `Aus RG/LF 3086294 vom 15.10.2026`。任一命中就停，见 ISSUES [PO 导入] D。
CREDIT_RE = re.compile(r"Gutschrift|Aus RG/LF \d+ vom")


#: 版式 B 表头；有它就走 parse_text_b
B_HEAD_RE = re.compile(r"^Menge Einheit Artikelbezeichnung PZN Symbol")
#: 版式 B 明细行：Menge 〈Einheit + 品名〉 PZN Symbol…(如 `SE N*`) Lief./Pos. [AVP] AEP Wert（AVP 可缺）
B_ROW_RE = re.compile(rf"^(\d+) (.+?) (\d{{8}}) ((?:[A-Z*]+ )*[A-Z*]+) \d+/\d+ (?:{_NUM} )?({_NUM}) ({_NUM})$")
#: 版式 B 页眉：单据类型 Kunden-Nr Beleg-Nr Datum（第 2 页 `NACHWEIS CHARGEN` 同单号，不算类型）
B_DOC_RE = re.compile(r"^([A-Z][A-Z ]*?) [\d.]+ (\d{6,8}) \d{2}\.\d{2}\.\d{4}\b")
#: 版式 B 底部：`Gesamtsumme EUR <7% 税基> <19% 税基> <总额>`、`<行数> Positionen <件数> Stück`（ü 常抽成乱码）
B_TOTAL_RE = re.compile(rf"^Gesamtsumme EUR ({_NUM}) ({_NUM}) ({_NUM})$")
B_COUNT_RE = re.compile(r"^(\d+) Positionen (\d+) St")

#: 版式 C（门户 Lieferschein 网页另存的 HTML）：明细表头前 6 列；Menge 形如 `2 STK`；金额后的 € 存成什么都不管，只取数字
C_HEAD = ("Menge", "Artikelbezeichnung", "PZN", "Ek.Preis", "Zeilenwert", "MwSt")
C_QTY_RE = re.compile(r"^(\d+) STK$")
C_NUM_RE = re.compile(rf"^({_NUM})\D*$")
C_TAX = {"7%": True, "19%": False}


def check_totals(lines, foot):
    """解析护栏（两个版式共用）：逐行 AEP × Menge = Gesamt；行数 / 件数 / 两档税基 / 总额与印的一致；
    `endsumme`（印的第二个总额）= total。"""
    bad = [l for l in lines if l["aep"] * l["qty"] != l["gesamt"]]
    if bad:
        raise StopError("以下行 AEP × Menge ≠ Gesamt，解析可能错位，不出表：\n" + "\n".join(
            f"  {l['pzn']} | {l['name']} | {l['aep']} × {l['qty']} ≠ {l['gesamt']}" for l in bad))
    got = dict(npos=len(lines), nqty=sum(l["qty"] for l in lines),
               full=sum((l["gesamt"] for l in lines if not l["reduced"]), Decimal(0)),
               reduced=sum((l["gesamt"] for l in lines if l["reduced"]), Decimal(0)),
               total=sum(l["gesamt"] for l in lines))
    diff = {k: (got[k], foot[k]) for k in got if got[k] != foot[k]}
    if foot["total"] != foot["endsumme"]:
        diff["endsumme"] = (foot["total"], foot["endsumme"])
    if diff:
        labels = dict(npos="行数", nqty="总件数", full="volle MwSt 税基", reduced="ermäß. MwSt 税基",
                      total="Warenwert gesamt", endsumme="Warenwert gesamt vs Endsumme")
        raise StopError("抽出的明细与 PDF 印的合计对不上（抽出 vs 印的），说明有行没抽到、抽错了或税码误读，不出表：\n"
                        + "\n".join(f"  {labels[k]}：{a} vs {b}" for k, (a, b) in diff.items()))


def parse_text_b(page_texts):
    """版式 B 各页文本 → (明细行列表, 单号)，行结构同 parse_text。单据类型不是 LIEFERSCHEIN / RECHNUNG 就停。"""
    lines, docs, kinds, total, count = [], set(), set(), None, None
    for i, text in enumerate(page_texts, 1):
        for ln in (text or "").splitlines():
            ln = ln.strip()
            m = B_ROW_RE.match(ln)
            if m:
                lines.append(dict(page=i, pzn=pzn8(m[3]), qty=int(m[1]), name=m[2], code=m[4],
                                  reduced="*" in m[4], aep=de_num(m[5]), gesamt=de_num(m[6])))
                continue
            m = B_DOC_RE.match(ln)
            if m:
                docs.add(m[2])
                if m[1] != "NACHWEIS CHARGEN":
                    kinds.add(m[1])
                continue
            m = B_TOTAL_RE.match(ln)
            if m:
                total = (de_num(m[1]), de_num(m[2]), de_num(m[3]))
                continue
            m = B_COUNT_RE.match(ln)
            if m:
                count = (int(m[1]), int(m[2]))
    if len(docs) != 1:
        raise StopError(f"页眉单号（Beleg-Nr.）认不出或不唯一：{sorted(docs)}")
    if not kinds or not kinds <= {"LIEFERSCHEIN", "RECHNUNG"}:
        raise StopError(f"单号 {next(iter(docs))} 的单据类型是 {sorted(kinds) or '认不出'}，不是 LIEFERSCHEIN / "
                        "RECHNUNG——若是贷项 / 退货请在 Odoo 走退货或供应商贷项，不做 PO 导入表。")
    if not lines:
        raise StopError("PDF 里没解析出任何明细行——版式不认识。")
    if total is None or count is None:
        raise StopError("PDF 里找不到底部合计（Gesamtsumme EUR / n Positionen）——版式不认识，不敢确定明细是否抽全。")
    check_totals(lines, dict(npos=count[0], nqty=count[1], reduced=total[0], full=total[1],
                             total=total[0] + total[1], endsumme=total[2]))
    return lines, docs.pop()


class _HtmlCells(HTMLParser):
    """HTML → (表格行列表（每行是各单元格文本）, 全部文本片段)。"""

    def __init__(self):
        super().__init__()
        self.rows, self.texts, self._cell = [], [], None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.rows.append([])
        elif tag in ("td", "th"):
            self._cell = ""

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None:
            if self.rows:
                self.rows[-1].append(" ".join(self._cell.split()))
            self._cell = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell += data
        if data.strip():
            self.texts.append(data.replace("\xa0", " ").strip())


def parse_html_c(html):
    """版式 C（门户 Lieferschein 网页 HTML）→ (明细行列表, 单号)，行结构同 parse_text。见 ISSUES [PO 导入] G：
    税直接看 MwSt 列（7% / 19%，别的值就停）；`Kz`（A / R）与税无关不读。页面底部只印 Gesamt (Netto)，
    护栏只有逐行 Menge × Ek.Preis = Zeilenwert 与 Σ = Gesamt (Netto)；同 PZN 多行照原样分行。"""
    p = _HtmlCells()
    p.feed(html)
    docs = [p.texts[i + 1] for i, t in enumerate(p.texts[:-1]) if t == "Lieferschein-Nr:"]
    if len(docs) != 1 or not docs[0].isdigit():
        raise StopError(f"页面上的 Lieferschein-Nr 认不出或不唯一：{docs}（一次只存一张 Lieferschein 的页面）")
    heads = [i for i, r in enumerate(p.rows) if tuple(r[:len(C_HEAD)]) == C_HEAD]
    if len(heads) != 1:
        raise StopError(f"页面里明细表头（{' '.join(C_HEAD)}）有 {len(heads)} 个——版式不认识。")
    lines, total, odd = [], None, []
    body = p.rows[heads[0] + 1:]
    for k, r in enumerate(body):
        if "Gesamt (Netto)" in r:
            nxt = [c for c in (body[k + 1] if k + 1 < len(body) else []) if c]
            m = C_NUM_RE.match(nxt[0]) if len(nxt) == 1 else None
            if not m:
                raise StopError(f"Gesamt (Netto) 下一行不是一个金额：{nxt}")
            total = de_num(m[1])
            break
        q = C_QTY_RE.match(r[0]) if len(r) >= len(C_HEAD) else None
        aep, gesamt = (C_NUM_RE.match(c) for c in r[3:5]) if q else (None, None)
        if not (q and r[2].isdigit() and aep and gesamt and r[5] in C_TAX):
            odd.append(r)
            continue
        lines.append(dict(page=1, pzn=pzn8(r[2]), qty=int(q[1]), name=r[1], code=r[5], reduced=C_TAX[r[5]],
                          aep=de_num(aep[1]), gesamt=de_num(gesamt[1])))
    if odd:
        raise StopError("明细表里有认不出的行（单位不是 STK / 税不是 7%·19% / 金额格式不对），不出表：\n"
                        + "\n".join(f"  {r}" for r in odd))
    if not lines:
        raise StopError("页面里没解析出任何明细行——版式不认识。")
    if total is None:
        raise StopError("页面里找不到 Gesamt (Netto)——不敢确定明细是否抽全。")
    bad = [l for l in lines if l["aep"] * l["qty"] != l["gesamt"]]
    if bad:
        raise StopError("以下行 Menge × Ek.Preis ≠ Zeilenwert，不出表：\n" + "\n".join(
            f"  {l['pzn']} | {l['name']} | {l['aep']} × {l['qty']} ≠ {l['gesamt']}" for l in bad))
    got = sum(l["gesamt"] for l in lines)
    if got != total:
        raise StopError(f"明细 Zeilenwert 合计 {got} ≠ 印的 Gesamt (Netto) {total}，不出表。")
    return lines, docs[0]


def read_source(path):
    """PDF → parse_text；门户另存的 .html / .htm → parse_html_c。"""
    if path.lower().endswith((".html", ".htm")):
        with open(path, encoding="utf-8", errors="replace") as f:
            return parse_html_c(f.read())
    return parse_text(read_pdf(path))


def refuse_credit_note(page_texts):
    """贷项单据（退量 / 退款）不是采购，按 PO 导入会把退货记成进货 → 停。"""
    hits = [ln.strip() for text in page_texts for ln in (text or "").splitlines() if CREDIT_RE.search(ln)]
    if hits:
        doc = next((m[1] for text in page_texts for ln in (text or "").splitlines()
                    if (m := DOC_RE.match(ln.strip()))), "?")
        raise StopError(f"单号 {doc} 是 Gutschrift（贷项 / 退货），不是采购，不出表——请在 Odoo 走退货或供应商贷项。命中：\n"
                        + "\n".join(f"  {h}" for h in hits))


def parse_text(page_texts):
    """各页文本 → (明细行列表, 单号)。行内含 pzn/name/qty/aep/gesamt/reduced/page。贷项单据先停。
    有版式 B 表头就转给 parse_text_b。"""
    if any(B_HEAD_RE.match(ln.strip()) for text in page_texts for ln in (text or "").splitlines()):
        return parse_text_b(page_texts)
    refuse_credit_note(page_texts)
    lines, docs, foot, endsumme = [], set(), None, None
    for i, text in enumerate(page_texts, 1):
        for ln in (text or "").splitlines():
            ln = ln.strip()
            m = ROW_RE.match(ln)
            if m:
                lines.append(dict(page=i, pzn=pzn8(m[1]), qty=int(m[2]), name=f"{m[3]} {m[4]} {m[5]}",
                                  code=m[6], reduced="M" in m[6].split(),
                                  aep=de_num(m[7]), gesamt=de_num(m[9])))
                continue
            m = ZERO_RE.match(ln)
            if m:
                words = m[3].split()
                name = "".join(words) if all(len(w) == 1 for w in words) else m[3]
                lines.append(dict(page=i, pzn=pzn8(m[1]), qty=int(m[2]), name=name, code="", reduced=False,
                                  aep=Decimal(0), gesamt=Decimal(0)))
                continue
            m = DOC_RE.match(ln)
            if m:
                docs.add(m[1])
                continue
            m = FOOT_RE.search(ln)
            if m:
                foot = dict(npos=int(m[1]), nqty=int(m[2]), full=de_num(m[3]),
                            reduced=de_num(m[4]), total=de_num(m[5]))
                continue
            m = END_RE.search(ln)
            if m:
                endsumme = de_num(m[1])
    if not lines:
        raise StopError("PDF 里没解析出任何明细行——版式不认识（本脚本只支持一家的 Rechnung/Lieferschein 版式）。")
    if len(docs) != 1:
        raise StopError(f"页眉单号认不出或不唯一：{sorted(docs)}")
    if foot is None or endsumme is None:
        raise StopError("PDF 里找不到底部合计（Anz. Pos. / Warenwert / Endsumme）——版式不认识，不敢确定明细是否抽全。")
    check_totals(lines, dict(foot, endsumme=endsumme))
    return lines, docs.pop()


def is_variant(code):
    """Internal Reference 带渠道 / 组合装尾缀（`_VO`、`_GW`、`xN`、`*N` 及组合）= 副本，采购不用。
    判定复用全仓 SKU 归一口径 `common/po._po_base_sku`（ISSUES [PO 导入] B，2026-10-20 用户定）。"""
    code = (code or "").strip()
    return _po_base_sku(code) != code


def match(lines, candidates, pick):
    """candidates: pzn → odoo_candidates 的元组。pick: pzn → 外部ID。

    采购只用基础品：先排除副本（`is_variant`），剩下恰好 1 个基础品才填进行里，被排除的副本记在
    `l["variants"]` 供核对表留痕；0 个或多个基础品收进 problems 原样返回，不猜、也不退回去用副本。
    """
    pick = {pzn8(k): v for k, v in pick.items()}
    unknown = set(pick) - {l["pzn"] for l in lines}
    if unknown:
        raise StopError(f"--pick 里的 PZN 不在（剔除后的）PDF 明细上：{sorted(unknown)}")
    problems = []
    for l in lines:
        hits = candidates.get(l["pzn"], [])
        base = [h for h in hits if not is_variant(h[2])]
        l["variants"] = [h for h in hits if is_variant(h[2])]
        if l["pzn"] in pick:
            if pick[l["pzn"]] in {h[0] for h in l["variants"]}:
                raise StopError(f"--pick {l['pzn']}={pick[l['pzn']]}：这是渠道 / 组合装副本，采购只用基础品。")
            chosen = [h for h in base if h[0] == pick[l["pzn"]]]
            if not chosen:
                raise StopError(f"--pick {l['pzn']}={pick[l['pzn']]}：这个外部 ID 不在该 PZN 的基础品命中里 "
                                f"{[h[0] for h in base]}")
            base = chosen
        if len(base) == 1:
            l["xmlid"], l["erp_name"], _, l["active"], _, l["erp_taxes"], l["dbid"] = base[0]
        else:
            problems.append((l, base))
    return problems


def _odoo():
    from odoo_api.odoo_client import Odoo
    return Odoo.connect(context={"active_test": False}, verbose=False)


def odoo_candidates(pzns):
    """Odoo 只读：pzn → [(__export__ 外部 ID, 品名, Internal Reference, active, 历史采购行数,
    默认进项税名称元组, 数据库 ID)]，含已归档品。"""
    od = _odoo()
    prods = od.execute("product.product", "search_read", [[["pzn", "in", sorted(pzns)]]],
                       {"fields": ["id", "name", "pzn", "default_code", "active", "supplier_taxes_id"]})
    ids = [p["id"] for p in prods]
    xids = od.execute("ir.model.data", "search_read",
                      [[["model", "=", "product.product"], ["module", "=", "__export__"], ["res_id", "in", ids]]],
                      {"fields": ["res_id", "name"]}) if prods else []
    by_res = {}
    for x in xids:
        by_res.setdefault(x["res_id"], []).append(f"__export__.{x['name']}")
    counts = {g["product_id"][0]: g["__count"] for g in od.execute(
        "purchase.order.line", "read_group", [[["product_id", "in", ids]]],
        {"fields": ["product_id"], "groupby": ["product_id"], "lazy": False})} if prods else {}
    tax_ids = sorted({t for p in prods for t in p["supplier_taxes_id"]})
    tax_names = {t["id"]: t["name"] for t in od.execute(
        "account.tax", "read", [tax_ids], {"fields": ["name"]})} if tax_ids else {}
    out = {}
    for p in prods:
        xs = by_res.get(p["id"], [])
        if len(xs) != 1:
            raise StopError(f"产品 {p['id']}（PZN {p['pzn']}，{p['name'].strip()}）的 __export__ 外部 ID 有 "
                            f"{len(xs)} 个，没法确定写哪个。先在 Odoo 里导出一次该产品生成外部 ID。")
        out.setdefault(pzn8(p["pzn"]), []).append(
            (xs[0], (p["name"] or "").strip(), p["default_code"] or "", p["active"], counts.get(p["id"], 0),
             tuple(sorted(tax_names[t] for t in p["supplier_taxes_id"])), p["id"]))
    return out


def odoo_existing(doc_no):
    """Odoo 只读：Vendor Reference = 单号的已有 PO → [(单号, 状态)]。"""
    return [(p["name"], p["state"]) for p in _odoo().execute(
        "purchase.order", "search_read", [[["partner_ref", "=", doc_no]]], {"fields": ["name", "state"]})]


def odoo_po_status(xmlids):
    """Odoo 只读：PO 外部 ID → (state, 订单行数, Vendor Reference)；Odoo 里找不到的不在结果里。"""
    od = _odoo()
    pairs = [x.split(".", 1) for x in xmlids]
    recs = od.execute("ir.model.data", "search_read",
                      [[["model", "=", "purchase.order"], ["module", "in", sorted({m for m, _ in pairs})],
                        ["name", "in", [n for _, n in pairs]]]], {"fields": ["module", "name", "res_id"]})
    by_res = {r["res_id"]: f"{r['module']}.{r['name']}" for r in recs}
    pos = od.execute("purchase.order", "read", [sorted(by_res)],
                     {"fields": ["state", "order_line", "partner_ref"]}) if by_res else []
    return {by_res[p["id"]]: (p["state"], len(p["order_line"]), p["partner_ref"] or "") for p in pos}


def blank_reason(status):
    """None = 空白草稿 PO（可随意写入）；否则返回不空白的原因。"""
    if status is None:
        return "Odoo 里找不到这张 PO"
    state, nlines, vref = status
    why = ([f"状态 {state}"] if state not in ("draft", "sent") else []) + \
          ([f"已有 {nlines} 行"] if nlines else []) + ([f"Vendor Reference 已是 {vref}"] if vref else [])
    return "、".join(why) or None


def _outputs_of(outdir, ref):
    """outdir 里这张 PO 已出过的导入表 → 对应的 PDF 名集合（去掉 unique_path 加的 `(n)` 序号）。"""
    prefix = f"{ref}_import_"
    names = os.listdir(outdir) if os.path.isdir(outdir) else []
    return {re.sub(r"\(\d+\)$", "", os.path.splitext(f)[0][len(prefix):])
            for f in names if f.startswith(prefix) and f.endswith(".xlsx")}


def choose_po(heads, colmap, outdir, stem, status_fn=odoo_po_status):
    """模板多张 PO、没给 --po 时自动选（ISSUES [PO 导入] F，2026-10-22 用户定）：空白 PO 之间没有区别，按模板行序取。

    同一 PDF 之前已出到某张仍空白的 PO → 沿用（重跑不换 PO）；输出目录里已为别的 PDF 出过表的 PO 视为已占用、跳过
    （分次跑同一模板的几张 PDF 不撞同一张）。没有可用的空白 PO 就停，列出每张为何不可用。
    """
    ref = colmap["ref"]
    status = status_fn([h[colmap["id"]] for h in heads])
    why = {h[ref]: blank_reason(status.get(h[colmap["id"]])) for h in heads}
    outs = {h[ref]: _outputs_of(outdir, h[ref]) for h in heads}
    same = [h for h in heads if stem in outs[h[ref]] and not why[h[ref]]]
    if same:
        print(f"自动分配：沿用 {same[0][ref]}（同一 PDF 之前已出到这张，仍空白）")
        return same[0]
    free = [h for h in heads if not why[h[ref]] and not outs[h[ref]]]
    if not free:
        raise StopError("模板里没有可自动写入的空白 PO，用 --po 指定或重新导出模板：\n" + "\n".join(
            f"  {h[ref]}：{why[h[ref]] or '本目录已为 ' + '、'.join(sorted(outs[h[ref]])) + ' 出过表'}" for h in heads))
    print(f"自动分配：{free[0][ref]}（空白 PO；模板里可用的空白 PO：{[h[ref] for h in free]}）")
    return free[0]


def read_template(template, po=None):
    """→ (workbook, 列名, 头信息列表[{列名: 值}], colmap{角色: 列名})。缺列或头信息不全就停。

    表头界面名（`Order Lines/Quantity`）与技术名（`order_line/product_qty`）都认，见 ALIASES。
    头信息取第 2 行起每个带 ID 的行（模板可含多张 PO）；给了 po（Order Reference）就只留那张。
    多张时选哪张由 run 调 choose_po 定。
    """
    wb = load_workbook(template)
    ws = wb.active
    cols = [c.value for c in ws[1]]
    colmap = {role: next((c for c in names if c and c in cols), None) for role, names in ALIASES.items()}

    def need(*roles):
        return " 或 ".join(f"{n}" for r in roles for n in ALIASES[r] if n)

    missing = [need(r) for r in ("id", "ref", "vendor", "price") if not colmap[r]]
    if not (colmap["pid"] or colmap["pdbid"]):
        missing.append(need("pid", "pdbid"))
    if not colmap["qty"]:
        # 2026-10-20 实测：Total Quantity 是 product_uom_qty，只读计算字段，导进去不改数量
        missing.append(need("qty") + (f"（模板里的 {colmap['total_qty']} 是只读字段、导入无效，导出时改勾 Quantity）"
                                      if colmap["total_qty"] else ""))
    if not (colmap["tax_id"] or colmap["tax"]):
        missing.append(need("tax_id", "tax"))
    if missing:
        raise StopError(f"模板缺列 {missing}；表头实际是 {cols}。在 Odoo 导出时把这些字段勾上重导。")
    heads = [{colmap[r]: ws.cell(row, cols.index(colmap[r]) + 1).value for r in ("id", "ref", "vendor")}
             for row in range(2, ws.max_row + 1) if ws.cell(row, cols.index(colmap["id"]) + 1).value]
    refs = [h[colmap["ref"]] for h in heads]
    if po:
        heads = [h for h in heads if h[colmap["ref"]] == po]
        if not heads:
            raise StopError(f"--po {po} 不在模板里；模板里的 PO：{refs}")
    bad = [h for h in heads if not all(h.values())] or ([] if heads else [{}])
    if bad:
        raise StopError(f"模板第 2 行起的 ID / Order Reference / Vendor 不全：{bad}")
    return wb, cols, heads, colmap


def build(wb, cols, head, colmap, lines, doc_no, out_path):
    """按模板列名列序写导入表：第一行带头信息（+ Vendor Reference），其余行只填明细列。

    产品列优先写外部 ID；模板只有 `.id` 列就写数据库 ID。税列有 `/id` 写外部 ID，否则写名称
    （两档名称在 Odoo 里唯一，2026-10-20 查实）。
    """
    ws = wb.active
    ws.delete_rows(2, ws.max_row)
    for k, l in enumerate(lines):
        tax_xmlid, tax_name = TAX_REDUCED if l["reduced"] else TAX_FULL
        vals = {colmap["qty"]: l["qty"], colmap["price"]: float(l["aep"])}
        if colmap["pid"]:
            vals[colmap["pid"]] = l["xmlid"]
        else:
            vals[colmap["pdbid"]] = l["dbid"]
        if colmap["tax_id"]:
            vals[colmap["tax_id"]] = tax_xmlid
        else:
            vals[colmap["tax"]] = tax_name
        if k == 0:
            vals.update(head)
            if colmap["vref"]:
                vals[colmap["vref"]] = doc_no
        ws.append([vals.get(c) for c in cols])
    wb.save(out_path)


def report(lines, dropped):
    print("\n核对表（PZN / PDF 名称 / 数量 / AEP / Gesamt / 税 / ERP 品名）")
    for k, l in enumerate(lines, 1):
        flag = "" if l["active"] else "  ⚠ 已归档"
        tax = TAX_REDUCED[1] if l["reduced"] else TAX_FULL[1]
        print(f"{k:>3} | {l['pzn']} | {l['name']} | {l['qty']} | {l['aep']} | {l['gesamt']} | {tax} | "
              f"{l['erp_name']}{flag}")
    off = [l for l in lines if l["erp_taxes"] != ((TAX_REDUCED if l["reduced"] else TAX_FULL)[1],)]
    if off:
        # 只提示不拦：表里写的是 PDF 税码对应的税；不一致说明产品主数据的默认进项税可能该改
        print("\n⚠ 以下产品在 Odoo 的默认进项税与 PDF 税码不一致（表里按 PDF 写），请核对产品主数据：")
        for l in off:
            want = (TAX_REDUCED if l["reduced"] else TAX_FULL)[1]
            print(f"    {l['pzn']} | {l['erp_name']} | PDF → {want} | 产品默认 {', '.join(l['erp_taxes']) or '无'}")
    skipped = [l for l in lines if l["variants"]]
    if skipped:
        print("\n自动排除的渠道 / 组合装副本（采购只用基础品）：")
        for l in skipped:
            print(f"    {l['pzn']} → 用 {l['xmlid']}；排除 {', '.join(h[2].strip() for h in l['variants'])}")
    if dropped:
        print("\n已剔除（建档后需手补进这张 PO）：")
        for l in dropped:
            print(f"    {l['pzn']} | {l['name']} | {l['qty']} | Gesamt {l['gesamt']}")
    r = sum((l["gesamt"] for l in lines if l["reduced"]), Decimal(0))
    f = sum((l["gesamt"] for l in lines if not l["reduced"]), Decimal(0))
    print(f"\n合计核对：出表 {len(lines)} 行 {sum(l['qty'] for l in lines)} 件；净额 {r + f}"
          f"（7% 税基 {r} + 19% 税基 {f}）；含税约 {(r * Decimal('1.07') + f * Decimal('1.19')).quantize(Decimal('0.01'))}"
          + (f"；剔除行合计 {sum(l['gesamt'] for l in dropped)}" if dropped else ""))


def run(pdf, template, outdir=None, exclude=(), pick=None, po=None,
        candidates_fn=odoo_candidates, existing_fn=odoo_existing, po_status_fn=odoo_po_status):
    wb, cols, heads, colmap = read_template(template, po)
    lines, doc_no = read_source(pdf)
    print(f"PDF 单号 {doc_no}，明细 {len(lines)} 行 / {sum(l['qty'] for l in lines)} 件"
          f"（页 {sorted({l['page'] for l in lines})}），合计 {sum(l['gesamt'] for l in lines)}")
    dup = existing_fn(doc_no)
    if dup:
        raise StopError(f"Odoo 里已有 Vendor Reference = {doc_no} 的采购单：{dup}——这张很可能已经录过，不出表。")
    lines, dropped = apply_exclude(lines, exclude)
    problems = match(lines, candidates_fn({l["pzn"] for l in lines}), pick or {})
    if problems:
        msg = ["以下行匹配异常，未出表（剔除用 --exclude PZN，多个基础品里选一个用 --pick PZN=外部ID）："]
        for l, base in problems:
            why = ("Odoo 里只有渠道 / 组合装副本、没有基础品——先建基础品，或这次剔掉" if not base and l["variants"]
                   else "Odoo 里查不到" if not base else f"{len(base)} 个基础品")
            msg.append(f"  {l['pzn']} | {l['name']} | 数量 {l['qty']} | {why}")
            msg += [f"      {h[0]}  {h[1]}  [{h[2]}]  历史采购 {h[4]} 行{'' if h[3] else '  已归档'}"
                    for h in base]
            msg += [f"      （已排除副本）{h[0]}  [{h[2]}]" for h in l["variants"]]
        raise StopError("\n".join(msg))
    outdir = outdir or os.path.join("output", f"{dt.date.today():%Y%m%d}")
    os.makedirs(outdir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(pdf))[0]
    head = heads[0] if len(heads) == 1 else choose_po(heads, colmap, outdir, stem, po_status_fn)
    out = unique_path(os.path.join(outdir, f"{head[colmap['ref']]}_import_{stem}.xlsx"))
    build(wb, cols, head, colmap, lines, doc_no, out)
    print(f"写出 {out}（{len(lines)} 行）；头信息 {head}"
          + (f"，Vendor Reference {doc_no}" if colmap["vref"] else "（模板无 Vendor Reference 列，单号没写进去）"))
    report(lines, dropped)
    return out


def main():
    ap = argparse.ArgumentParser(description="供应商 Rechnung/Lieferschein PDF → Odoo 采购单导入表")
    ap.add_argument("pdf", help="供应商 Rechnung/Lieferschein PDF，或门户 Lieferschein 网页另存的 .html")
    ap.add_argument("template", help="Odoo 导出的 purchase.order 模板 xlsx（第 2 行起带 ID / Order Reference / Vendor）")
    ap.add_argument("--po", metavar="PO号", help="模板含多张 PO 时指定写进哪张（Order Reference，如 P12380）；不给就自动选空白 PO")
    ap.add_argument("-o", "--outdir", help="输出目录（默认 output/YYYYMMDD）")
    ap.add_argument("--exclude", nargs="+", default=[], metavar="PZN", help="剔除这些 PZN 的行（用户裁决后用）")
    ap.add_argument("--pick", nargs="+", default=[], metavar="PZN=外部ID",
                    help="一个 PZN 命中多个产品时指定用哪个（用户裁决后用）")
    a = ap.parse_args()
    pick = {}
    for kv in a.pick:
        if "=" not in kv:
            ap.error(f"--pick 要写成 PZN=外部ID：{kv}")
        k, v = kv.split("=", 1)
        pick[k.strip()] = v.strip()
    try:
        run(a.pdf, a.template, a.outdir, a.exclude, pick, a.po)
    except StopError as e:
        print(f"\n✗ {e}", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
