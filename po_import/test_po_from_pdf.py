"""po_from_pdf 测试：PDF 文本解析与四道护栏、PZN 匹配裁决、防重、按模板写表。

不连 Odoo、不读真 PDF：解析吃文本、匹配吃假候选表，模板用代码造（*.xlsx 不进仓库）。
文本行照 2026-10-20 实测 Rechnung/Lieferschein 的 pdfplumber 抽取结果构造（€ 抽成乱码），品名与价格是造的。
"""
import os
import sys
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from po_import import po_from_pdf as P  # noqa: E402

ROWS = [
    "17173992 2 120ST KAP PRODUKT A S M 10,00 0,00 20,00",
    "02758221 5 100ST HKP PRODUKT B RETARD P S 4,50 9,95 22,50",
    "20101140 10 15ML OEL PRODUKT C S 1,25 0,00 12,50",
]
FOOT = "RX-DPP:1,00 � 3/17 35,00 � 20,00 � 55,00 �\n**** Endsumme **** 55,00 �"
PAGE = "1234567 1\n" + "\n".join(ROWS) + "\n" + FOOT + "\nKUNDE GmbH\n7654321\n20.10.2026"

COLS = ["ID", "Order Reference", "Vendor", P.COL_VREF, P.COL_PID, P.COL_QTY, P.COL_PRICE, P.COL_TAX_ID]


def test_parse_lines_doc_and_tax_class():
    lines, doc = P.parse_text([PAGE])
    assert doc == "1234567"
    assert [l["pzn"] for l in lines] == ["17173992", "02758221", "20101140"]
    assert [l["reduced"] for l in lines] == [True, False, False]
    # 品名末尾的 P 不是税码字母，留在品名里
    assert lines[1]["name"] == "100ST HKP PRODUKT B RETARD P" and lines[1]["code"] == "S"
    assert lines[0]["aep"] == Decimal("10.00") and lines[2]["qty"] == 10


@pytest.mark.parametrize("bad", [
    PAGE.replace("3/17", "4/17"),                       # 行数对不上 = 漏行
    PAGE.replace("3/17", "3/18"),                       # 件数对不上
    PAGE.replace("PRODUKT A S M", "PRODUKT A S"),       # 税码误读 → 两档税基对不上
    PAGE.replace("**** Endsumme **** 55,00", "**** Endsumme **** 56,00"),
    PAGE.replace("4,50 9,95 22,50", "4,50 9,95 22,60"),  # 行内 AEP×Menge ≠ Gesamt
    PAGE.replace(FOOT, ""),                              # 没有底部合计
    PAGE.replace("1234567 1\n", ""),                     # 没有单号
    "nichts hier",
])
def test_parse_refuses(bad):
    with pytest.raises(P.StopError):
        P.parse_text([bad])


def _lines():
    return P.parse_text([PAGE])[0]


def _cand(xid, n=0, taxes=(P.TAX_FULL[1],), dbid=0, ref=None):
    return (xid, "ERP " + xid, ref or "Base_" + xid.split(".")[-1], True, n, taxes, dbid)


def test_match_unique_none_multi_and_pick():
    lines = _lines()
    cands = {"17173992": [_cand("__export__.a")],
             "02758221": [_cand("__export__.b", 75), _cand("__export__.b2")]}
    problems = P.match(lines, cands, {})
    assert lines[0]["xmlid"] == "__export__.a"
    assert [(l["pzn"], len(h)) for l, h in problems] == [("02758221", 2), ("20101140", 0)]

    lines = _lines()
    problems = P.match(lines, cands, {"2758221": "__export__.b"})
    assert lines[1]["xmlid"] == "__export__.b"
    assert [l["pzn"] for l, _ in problems] == ["20101140"]
    with pytest.raises(P.StopError):
        P.match(_lines(), cands, {"02758221": "__export__.nope"})


def _template(tmp_path, cols=COLS):
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    ws.append(["__export__.purchase_order_1", "P00001", "Vendor X"] + [None] * (len(cols) - 3))
    p = tmp_path / "purchase.order.xlsx"
    wb.save(p)
    return str(p)


def _run(tmp_path, monkeypatch, template, existing=(), **kw):
    monkeypatch.setattr(P, "read_pdf", lambda _: [PAGE])
    cands = {"17173992": [_cand("__export__.a", dbid=11)], "02758221": [_cand("__export__.b", dbid=12)],
             "20101140": [_cand("__export__.c", dbid=13)]}
    return P.run("L 1.pdf", template, str(tmp_path / "out"), candidates_fn=lambda _: cands,
                 existing_fn=lambda _: list(existing), **kw)


def test_run_writes_template(tmp_path, monkeypatch):
    out = _run(tmp_path, monkeypatch, _template(tmp_path))
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert list(rows[0]) == COLS
    assert rows[1] == ("__export__.purchase_order_1", "P00001", "Vendor X", "1234567",
                       "__export__.a", 2, 10.0, P.TAX_REDUCED[0])
    assert rows[2] == (None, None, None, None, "__export__.b", 5, 4.5, P.TAX_FULL[0])
    assert len(rows) == 4


def test_run_tax_by_name_and_exclude(tmp_path, monkeypatch):
    cols = ["ID", "Order Reference", "Vendor", P.COL_PID, P.COL_QTY, P.COL_PRICE, P.COL_TAX]
    out = _run(tmp_path, monkeypatch, _template(tmp_path, cols), exclude=["20101140"])
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert [r[-1] for r in rows[1:]] == [P.TAX_REDUCED[1], P.TAX_FULL[1]]


def test_run_refuses_duplicate(tmp_path, monkeypatch):
    with pytest.raises(P.StopError, match="已有"):
        _run(tmp_path, monkeypatch, _template(tmp_path), existing=[("P12345", "done")])
    assert not (tmp_path / "out").exists()


def test_template_missing_tax_column(tmp_path):
    with pytest.raises(P.StopError, match="Taxes"):
        P.read_template(_template(tmp_path, COLS[:-1]))


def test_template_total_quantity_is_explained(tmp_path):
    cols = [P.COL_TOTAL_QTY if c == P.COL_QTY else c for c in COLS]
    with pytest.raises(P.StopError, match="只读"):
        P.read_template(_template(tmp_path, cols))


def test_report_flags_default_tax_mismatch(tmp_path, monkeypatch, capsys):
    # 17173992 在 PDF 是 M（7%），候选默认税给的是 19% → 提示；其余两行一致 → 不提示
    _run(tmp_path, monkeypatch, _template(tmp_path))
    tail = capsys.readouterr().out.split("不一致")[1].split("合计核对")[0]
    assert "17173992" in tail and "02758221" not in tail and "20101140" not in tail


def test_run_technical_names_with_db_id(tmp_path, monkeypatch):
    # 2026-10-20 用户第 28 份模板：技术名表头，产品列是数据库 ID（.id），税列按名称
    cols = ["id", "partner_ref", "name", "partner_id", "order_line/product_id/.id",
            "order_line/product_qty", "order_line/price_unit", "order_line/taxes_id"]
    wb = Workbook()
    ws = wb.active
    ws.append(cols)
    ws.append(["__export__.purchase_order_1", None, "P00001", "Vendor X", False, None, None, None])
    p = tmp_path / "purchase.order.xlsx"
    wb.save(p)
    out = _run(tmp_path, monkeypatch, str(p))
    assert os.path.basename(out).startswith("P00001_import_")
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert rows[1] == ("__export__.purchase_order_1", "1234567", "P00001", "Vendor X", 11, 2, 10.0,
                       P.TAX_REDUCED[1])
    assert rows[2] == (None, None, None, None, 12, 5, 4.5, P.TAX_FULL[1])


def test_template_prefers_external_product_id(tmp_path, monkeypatch):
    cols = COLS + ["order_line/product_id/.id"]
    rows = list(load_workbook(_run(tmp_path, monkeypatch, _template(tmp_path, cols))).active
                .iter_rows(values_only=True))
    assert rows[1][4] == "__export__.a" and rows[1][-1] is None


@pytest.mark.parametrize("code, variant", [
    ("Orthomol_17209603", False), ("Orthomol_17209603_GW", True), ("Legalon_11548215_VO", True),
    ("Orthomol_17209603x2_GW", True), ("Avene_00670485x2", True), ("Babix_04459669x3_VO", True),
    ("Betaisodona_00721478 ", False), ("", False),
])
def test_is_variant(code, variant):
    assert P.is_variant(code) is variant


def test_match_drops_variants():
    cands = {
        # 基础品 + 两个副本 → 自动用基础品
        "17173992": [_cand("__export__.a", ref="A_17173992"), _cand("__export__.a_gw", ref="A_17173992_GW"),
                     _cand("__export__.a2", ref="A_17173992x2_GW")],
        # 只有副本 → 停，不退回去用副本
        "02758221": [_cand("__export__.b_vo", ref="B_02758221_VO"), _cand("__export__.b_gw", ref="B_02758221_GW")],
        # 两个基础品 + 一个副本 → 停，只列基础品
        "20101140": [_cand("__export__.c", ref="C_20101140"), _cand("__export__.c_old", ref="C_123456"),
                     _cand("__export__.c_vo", ref="C_20101140_VO")],
    }
    lines = _lines()
    problems = P.match(lines, cands, {})
    assert lines[0]["xmlid"] == "__export__.a"
    assert [h[0] for h in lines[0]["variants"]] == ["__export__.a_gw", "__export__.a2"]
    assert [(l["pzn"], [h[0] for h in b]) for l, b in problems] == [
        ("02758221", []), ("20101140", ["__export__.c", "__export__.c_old"])]

    lines = _lines()
    problems = P.match(lines, cands, {"20101140": "__export__.c_old"})
    assert lines[2]["xmlid"] == "__export__.c_old" and [l["pzn"] for l, _ in problems] == ["02758221"]
    with pytest.raises(P.StopError, match="副本"):
        P.match(_lines(), cands, {"02758221": "__export__.b_vo"})


def test_parse_zero_value_insert_row():
    # 0 元附件行（ISSUES [PO 导入] C）：品名逐字母带空格、无 Code，照样进明细、单价 0、按 19%
    page = PAGE.replace(ROWS[0], "05855555 1 B E I P A C K 0,00 0,00 0,00\n" + ROWS[0]).replace("3/17", "4/18")
    lines, _ = P.parse_text([page])
    z = lines[0]
    assert (z["pzn"], z["qty"], z["name"], z["aep"], z["gesamt"], z["reduced"]) == \
        ("05855555", 1, "BEIPACK", Decimal(0), Decimal(0), False)
    # 金额非 0 的同形行不归它管 → 漏行被护栏拦
    with pytest.raises(P.StopError):
        P.parse_text([page.replace("B E I P A C K 0,00 0,00 0,00", "B E I P A C K 1,00 0,00 1,00")])


@pytest.mark.parametrize("extra", [
    "Gutschrift Menge + Wert",                    # 页尾单据类型（ISSUES [PO 导入] D）
    "Aus RG/LF 7654320 vom 15.10.2026",           # 只有引用原单的行也算贷项
])
def test_parse_refuses_credit_note(extra):
    with pytest.raises(P.StopError, match="Gutschrift"):
        P.parse_text([PAGE + "\n" + extra])
    assert P.parse_text([PAGE + "\nNormal-Auftrag"])[1] == "1234567"


# 版式 B（ISSUES [PO 导入] E）：照 2026-10-22 实测 Lieferschein 的 pdfplumber 抽取结果构造，品名与价格是造的。
B_ROWS = [
    "2 30 ST PRODUKT D KAPSEL 01534672 SE N 2/6 3,45 2,20 4,40",
    "20 120 ST PRODUKT E SY 07625016 SE N* 4/13 36,95 21,19 423,80",
    "1 30X3 G PRODUKT F PUL 17458393 S N* 3/8 26,43 26,43",          # AVP 缺
    "* 30 BESTELLT, 21 ZU AKTIONSKOND. GELIEF.",                    # 说明行，忽略
    "5 N2 20 ST PRODUKT G TAB 00624605 S N 3/9 9,97 4,40 22,00",
]
B_PAGE = ("Kunden-Nr. Beleg-Nr. Datum Blatt\nLIEFERSCHEIN 19.1523 226082 21.10.2026 1 / 1\n"
          "Menge Einheit Artikelbezeichnung PZN Symbol Lief./Pos. AVP AEP Wert\n" + "\n".join(B_ROWS)
          + "\n7% USt. 19% USt.\nGesamtsumme EUR 450,23 26,40 476,63\n4 Positionen 28 St�ckZeit:19.11")
B_CHARGEN = ("NACHWEIS CHARGEN 19.1523 226082 21.10.2026 1 / 1\n"
             "2 30 ST PRODUKT D KAPSEL 01534672 261193")


def test_parse_b_lines_doc_and_tax_class():
    lines, doc = P.parse_text([B_PAGE, B_CHARGEN])
    assert doc == "226082"
    assert [l["pzn"] for l in lines] == ["01534672", "07625016", "17458393", "00624605"]
    assert [l["reduced"] for l in lines] == [False, True, True, False]
    assert [l["qty"] for l in lines] == [2, 20, 1, 5]
    assert lines[2]["aep"] == Decimal("26.43") and lines[3]["name"] == "N2 20 ST PRODUKT G TAB"


@pytest.mark.parametrize("bad", [
    B_PAGE.replace("4 Positionen", "5 Positionen"),                       # 漏行
    B_PAGE.replace("28 St", "29 St"),                                     # 件数
    B_PAGE.replace("SE N* 4/13", "SE N 4/13"),                            # 税档误读
    B_PAGE.replace("476,63", "476,64"),                                   # 总额
    B_PAGE.replace("2,20 4,40", "2,20 4,50"),                             # AEP×Menge≠Wert
    B_PAGE.replace("Gesamtsumme EUR 450,23 26,40 476,63", ""),            # 无合计
    B_PAGE.replace("LIEFERSCHEIN 19.1523", "GUTSCHRIFT 19.1523"),         # 单据类型不对
    B_PAGE.replace("LIEFERSCHEIN 19.1523 226082 21.10.2026 1 / 1\n", ""),  # 无单号
])
def test_parse_b_refuses(bad):
    with pytest.raises(P.StopError):
        P.parse_text([bad])


def _two_po_template(tmp_path):
    p = _template(tmp_path)
    wb = load_workbook(p)
    wb.active.append(["__export__.purchase_order_2", "P00002", "Vendor X"] + [None] * (len(COLS) - 3))
    wb.save(p)
    return p


BLANK = ("draft", 0, "")


def test_template_multiple_pos_and_explicit_po(tmp_path, monkeypatch):
    p = _two_po_template(tmp_path)
    assert [h["Order Reference"] for h in P.read_template(p)[2]] == ["P00001", "P00002"]
    with pytest.raises(P.StopError, match="P00009"):
        P.read_template(p, "P00009")
    # 给了 --po 就用那张，不查空白
    out = _run(tmp_path, monkeypatch, p, po="P00002", po_status_fn=lambda _: pytest.fail("不该查"))
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert rows[1][:4] == ("__export__.purchase_order_2", "P00002", "Vendor X", "1234567")
    assert len(rows) == 4 and os.path.basename(out).startswith("P00002_")


@pytest.mark.parametrize("status", [("purchase", 0, ""), ("draft", 3, ""), ("draft", 0, "999")])
def test_blank_reason(status):
    assert P.blank_reason(BLANK) is None
    assert P.blank_reason(status) and P.blank_reason(None)


def test_auto_po_picks_blank_skips_taken_and_reuses_same_pdf(tmp_path, monkeypatch):
    p = _two_po_template(tmp_path)
    both = {"__export__.purchase_order_1": BLANK, "__export__.purchase_order_2": BLANK}
    # 两张都空白 → 按行序取第一张
    out1 = _run(tmp_path, monkeypatch, p, po_status_fn=lambda _: both)
    assert os.path.basename(out1) == "P00001_import_L 1.xlsx"
    # 同一 PDF 重跑 → 沿用 P00001（unique_path 加序号）
    out2 = _run(tmp_path, monkeypatch, p, po_status_fn=lambda _: both)
    assert os.path.basename(out2) == "P00001_import_L 1(1).xlsx"
    # 另一张 PDF → P00001 已被占用，去 P00002
    monkeypatch.setattr(P, "read_pdf", lambda _: [PAGE])
    cands = {"17173992": [_cand("__export__.a")], "02758221": [_cand("__export__.b")],
             "20101140": [_cand("__export__.c")]}
    out3 = P.run("L 2.pdf", p, str(tmp_path / "out"), candidates_fn=lambda _: cands,
                 existing_fn=lambda _: [], po_status_fn=lambda _: both)
    assert os.path.basename(out3) == "P00002_import_L 2.xlsx"


def test_auto_po_skips_non_blank_and_stops_when_none(tmp_path, monkeypatch):
    p = _two_po_template(tmp_path)
    st = {"__export__.purchase_order_1": ("draft", 2, ""), "__export__.purchase_order_2": BLANK}
    out = _run(tmp_path, monkeypatch, p, po_status_fn=lambda _: st)
    assert os.path.basename(out).startswith("P00002_")
    st["__export__.purchase_order_2"] = ("purchase", 0, "")
    with pytest.raises(P.StopError, match="已有 2 行"):
        _run(tmp_path, monkeypatch, p, po_status_fn=lambda _: st)
