"""lot_entry 测试：解析小件 + 「导出 → 录入表 → 操作员填写 → 导入表」全程。

全程用代码造的假导出（*.xlsx 一律 gitignore，真实导出不进仓库），
列名用英文标签版，与 2026-10-11 用户实际导出 `stock.picking (9).xlsx` 同构。
"""
import os
import re
import sys
from collections import OrderedDict
from datetime import date, datetime

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import lot_entry as L  # noqa: E402


# ---------------------------------------------------------------- 解析小件
@pytest.mark.parametrize("raw,want", [
    ("11.2027", date(2027, 11, 30)),
    ("2.2028", date(2028, 2, 29)),          # 闰年月底
    ("04/2029", date(2029, 4, 30)),
    ("30.01.2028", date(2028, 1, 30)),
    ("2027-10", date(2027, 10, 31)),
    ("2027-10-05", date(2027, 10, 5)),
    ("05.28", date(2028, 5, 31)),           # 操作员习惯 MM.YY
    ("5/28", date(2028, 5, 31)),
    ("02.28", date(2028, 2, 29)),
    ("31.10.29", date(2029, 10, 31)),       # DD.MM.YY 照抄包装全日期（R6）
    ("1.6.28", date(2028, 6, 1)),
    ("27-11-28", date(2028, 11, 27)),
    (" 12 . 2027 ", date(2027, 12, 31)),
    (datetime(2028, 3, 1), date(2028, 3, 1)),
])
def test_parse_mhd_ok(raw, want):
    assert L.parse_mhd(raw) == (want, None)


@pytest.mark.parametrize("raw", ["28.05", "13.28", "13.2027", "31.02.2028", "31.02.28", "2027", "11,2027x", 11.2027])
def test_parse_mhd_rejects(raw):
    d, err = L.parse_mhd(raw)
    assert d is None and err


def test_parse_mhd_yy_mm_hint():
    d, err = L.parse_mhd("28.05")
    assert d is None and "05.28" in err


def test_mhd_too_far_ahead_is_error():
    b = dict(row=3, name="X", soll=1, entries=[(3, "AB1234", "05.82", None)])
    lots, issues = L.check_block(b)
    assert any(lvl == "ERROR" and "years ahead" in m for lvl, _, m in issues)


def test_mhd_dd_mm_yy_too_far_ahead_is_error():
    b = dict(row=3, name="X", soll=1, entries=[(3, "AB1234", "15.09.82", None)])
    lots, issues = L.check_block(b)
    assert any(lvl == "ERROR" and "years ahead" in m for lvl, _, m in issues)


def test_parse_mhd_empty():
    assert L.parse_mhd(None) == (None, None)
    assert L.parse_mhd("  ") == (None, None)


def test_name_mhd_months():
    assert L.name_mhd_months("Mykoderm Mund-Gel 20g(MHD 2027/10)") == {(2027, 10)}
    assert L.name_mhd_months("Bepanthen Augentropfen MHD2029/04") == {(2029, 4)}
    assert L.name_mhd_months("Otriven 10ml(MHD05.28)") == {(2028, 5)}
    assert L.name_mhd_months("Crataegutt 50StMHD28.07") == {(2028, 7)}
    assert L.name_mhd_months("Elmex Kariesschutz 400ml") == set()


def test_new_pzn():
    assert L.new_pzn("Doppelherz Knoblauch 480St. (New PZN 19765739)") == "19765739"
    assert L.new_pzn("Foo (neue PZN: 1234567)") == "1234567"
    assert L.new_pzn("Rosuvastatin Hexal 10 Mg 100 St") is None


def test_pzn_rich():
    plain = str(L.pzn_rich("DOPH_15994609", "Doppelherz 480St."))
    assert plain == "DOPH_15994609"
    r = L.pzn_rich("DOPH_15994609", "Doppelherz 480St. (New PZN 19765739)")
    assert str(r).startswith("19765739") and "alt: DOPH_15994609" in str(r)
    assert r[0].font.b                                           # 新 PZN 加粗打头


def test_split_product_flattens_newlines():
    assert L.split_product("[DOPH_12454290] Doppelherz \nFlohsamenschalen \n250 G") == \
        ("DOPH_12454290", "Doppelherz Flohsamenschalen 250 G")


# ---------------------------------------------------------------- 全程
HEADER = ["Operations/ID", "Operations/Stock Move/ID", "Operations/Product",
          "Operations/Lot/Serial Number Name", "Operations/Done", "Operations/Unit of Measure",
          "Operations/From", "Operations/To", "Operations/Expiration Date", "ID",
          "External ID", "Contact", "Operations/Stock Move/Demand", "Source Document",
          "Reference"]


def X(s):
    """测试用 External ID：真实导出里是 `__export__.stock_move_line_601322_7a578f8e` 形态。"""
    return "__export__." + s


def _export(path):
    rows = [
        # 单 1：一个产品
        ("ml1", "mv1", "[Rosu_1] Rosuvastatin 100 St", 1000, "pk1", "Hexal", "P1", "WH/IN/1"),
        # 单 2：多个产品
        ("ml2", "mv2", "[Dopp_2] Doppelherz 480St", 360, "pk2", "Queisser", "P2", "WH/IN/2"),
        ("ml3", "mv3", "[Wolz_3] Cardio Care", 2, None, None, None, None),
        ("ml4", "mv4", "[Derm_4] Badesalz 5Kg", 2, None, None, None, None),
        ("ml5", "mv5", "[Flo_5] Flohsamen MHD 2028/05", 1, None, None, None, None),
        ("ml6", "mv6", "[Ome_6] Omega", 2, None, None, None, None),
        ("ml7", "mv7", "[Hylo_7] Hylo Dual", 9, None, None, None, None),
        ("ml8", "mv8", "[Elm_8] Elmex", 2, None, None, None, None),
        ("ml9", "mv9", "[Untouched_9] Nobody fills me", 1, None, None, None, None),
        # 单 3：没人录 → 整单不进导入表
        ("ml10", "mv10", "[Empty_10] Empty", 5, "pk3", "X", "P3", "WH/IN/3"),
    ]
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    for ml, mv, prod, dem, pk, partner, origin, ref in rows:
        ws.append([X(ml), X(mv), prod, None, 0, "Units", "Partner Locations/Vendors", "WH/Stock",
                   None, "123" if pk else None, pk and X(pk), partner, dem, origin, ref])
    wb.save(path)


def _fill(path):
    """模拟操作员：按产品编码找到产品行，往下面的录入行填。"""
    wb = load_workbook(path)
    ws = wb["Erfassung"]
    def locate():
        return {ws.cell(r, L.C_TECH0 + L.TECH.index("code")).value: r for r in range(2, ws.max_row + 1)
                if ws.cell(r, L.C_TECH0).value == "P"}
    prod = locate()

    def put(code, k, lot=None, mhd=None, qty=None):
        r = locate()[code] + k          # 第 1 个批次就在产品行上
        for col, v in ((L.C_LOT, lot), (L.C_MHD, mhd), (L.C_QTY, qty)):
            if v is not None:
                ws.cell(r, col, v)

    put("Rosu_1", 0, "25P039A", "11.2027")                      # 单批次不写数量 → 取应收
    put("Dopp_2", 0, "AB1234", "30.01.2028", 300)
    put("Dopp_2", 1, "4X3X-A", "04.2029", 60)
    put("Wolz_3", 0, "WZ0001", "12.2027", 1)                    # 合计 1 ≠ 2 → 警告
    put("Derm_4", 0, mhd="06.2028")                             # 无批次只写 MHD → 警告、跳过
    put("Flo_5", 0, "FL5555", "07.2028")                        # 与品名 MHD 2028/05 不符 → 警告
    put("Ome_6", 0, "OM1111", "28.05", 1)                       # YY.MM → 错误
    put("Ome_6", 1, "OM2222", "06.2028", 1)
    # Hylo 三批：操作员在续行下面（紧贴下一产品的上方）插一行
    ws.insert_rows(prod["Hylo_7"] + L.ENTRY_ROWS)
    prod = locate()
    put("Hylo_7", 0, "HY0001", "01.2028", 3)
    put("Hylo_7", 1, "HY0002", "02.2028", 3)
    put("Hylo_7", 2, "HY0003", "03.2028", 3)
    put("Elm_8", 0, "EL 12", "01.2029")                         # 批次带空格 → 错误
    wb.save(path)


def test_end_to_end(tmp_path):
    exp = tmp_path / "export.xlsx"
    _export(exp)
    kind, files, _ = L.run(str(exp), str(tmp_path))
    assert kind == "entry"
    entry = files[0]

    ws = load_workbook(entry)["Erfassung"]
    heads = [r for r in range(2, ws.max_row + 1) if ws.cell(r, L.C_TECH0).value == "H"]
    # R10：三家各一组；Queisser 8 个产品排前，Hexal / X 都是 1 个、按名称；换供应商前分页
    assert [ws.cell(r, 1).value.split(" · ")[0] for r in heads] == ["WH/IN/2", "WH/IN/1", "WH/IN/3"]
    assert [b.id for b in ws.row_breaks.brk] == [r - 1 for r in heads[1:]]
    p = heads[1] + 1
    assert ws.cell(p, L.C_TECH0 + L.TECH.index("code")).value == "Rosu_1" and ws.cell(p, L.C_SOLL).value == 1000
    assert ws.cell(p, L.C_NAME).value == "Rosuvastatin 100 St"
    for r in (p, p + 1):                                         # 产品行本身 + 续行都可录
        assert ws.cell(r, L.C_LOT).number_format == "@"
        assert ws.cell(r, L.C_LOT).protection.locked is False
    assert ws.cell(p, L.C_NAME).protection.locked is not False
    assert ws.protection.sheet is True
    assert [ws.cell(1, c).value for c in (L.C_LOT, L.C_STAT)] == ["LOT", "Status"]
    assert ws.print_area == "'Erfassung'!$A$1:$F$%d" % (ws.max_row - 1)  # Σ / Status 不打印
    assert ws.cell(3, L.C_STAT).value.startswith("=IF(")
    # 公式行锚点只许产品行本身与 END：写死 $I{row+1} 时，选中第 2 行插行会把它推下去、漏算插入行（R9）
    end = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, L.C_TECH0).value == "END")
    for r in range(2, end):
        if ws.cell(r, L.C_TECH0).value == "P":
            anchors = {int(n) for n in re.findall(r"[A-Z]\$?(\d+)", ws.cell(r, L.C_STAT).value)}
            assert anchors == {r, end}, (r, anchors)
    # 单内按品名字母序（导出里是 Doppelherz, Cardio, Badesalz, Flohsamen, Omega, Hylo, Elmex, Nobody）
    groups = []
    for r in range(2, ws.max_row + 1):
        t = ws.cell(r, L.C_TECH0).value
        if t == "H":
            groups.append([])
        elif t == "P":
            groups[-1].append(ws.cell(r, L.C_NAME).value)
    assert len(groups[0]) == 8
    assert groups[0] == sorted(groups[0], key=str.casefold)

    _fill(entry)
    kind, files, summary = L.run(entry, str(tmp_path))
    assert kind == "import"
    imp, chk = files

    rows = list(load_workbook(imp).active.iter_rows(values_only=True))
    assert list(rows[0]) == L.IMPORT_HEADER
    body = rows[1:]
    by_lot = {r[4]: r for r in body if r[4]}
    pickings = [r[0] for r in body if r[0]]
    assert pickings == [X("pk2"), X("pk1")]                            # 顺序随录入表（R10）；单 3 没人录 → 不出现

    assert by_lot["25P039A"][1] == X("ml1") and by_lot["25P039A"][5] == 1000
    assert by_lot["25P039A"][9] == datetime(2027, 11, 30)
    assert by_lot["AB1234"][1] == X("ml2") and by_lot["AB1234"][5] == 300
    assert by_lot["4X3X-A"][1] is None and by_lot["4X3X-A"][2] == X("mv2")
    assert by_lot["4X3X-A"][3] == "[Dopp_2] Doppelherz 480St"
    assert by_lot["4X3X-A"][9] == datetime(2029, 4, 30)
    assert [by_lot[x][1] for x in ("HY0001", "HY0002", "HY0003")] == [X("ml7"), None, None]
    assert "WZ0001" in by_lot and "FL5555" in by_lot             # 只警告的照写
    assert not ({"OM1111", "OM2222", "EL 12"} & set(by_lot))    # 标红的整产品不写

    # 没写批次的已有行原样带上（防 Odoo 把未列出的行替换掉）
    lines = [r[1] for r in body if r[1]]
    for ml in (X("ml4"), X("ml6"), X("ml8"), X("ml9")):
        assert ml in lines
    assert all(r[4] is None and r[5] == 0 for r in body if r[1] in (X("ml4"), X("ml6"), X("ml8"), X("ml9")))
    assert X("ml10") not in lines

    check = list(load_workbook(chk).active.iter_rows(values_only=True))[1:]
    msgs = {(c[0], c[2]) for c in check}
    assert ("WARN", "Wolz_3") in msgs
    assert ("WARN", "Derm_4") in msgs and ("NO LOT", "Derm_4") in msgs
    assert ("WARN", "Flo_5") in msgs
    assert ("ERROR", "Ome_6") in msgs and ("SKIP", "Ome_6") in msgs
    assert ("ERROR", "Elm_8") in msgs
    assert ("NO LOT", "Untouched_9") in msgs


def _pk(partner, n, ref):
    return {"ref": ref, "partner": partner, "origin": None, "moves": {i: None for i in range(n)}}


def test_group_by_vendor():
    """R10：四家固定在前（按名开头单词、不分大小写）→ 其余按组合计产品数 → 无 Contact 最后；组内按产品数、同数按 WH/IN 号。"""
    pickings = OrderedDict([
        ("a", _pk("Hexal AG (Sandoz)", 3, "WH/IN/20")),
        ("b", _pk("AEP GmbH", 9, "WH/IN/21")),
        ("c", _pk("Queisser Pharma", 2, "WH/IN/22")),
        ("d", _pk("PHOENIX Pharmahandel GmbH & Co KG", 2, "WH/IN/30")),
        ("e", _pk(None, 50, "WH/IN/23")),
        ("f", _pk("Queisser Pharma", 2, "WH/IN/24")),
        ("g", _pk("Otto Geilenkirchen GmbH & Co. KG", 1, "WH/IN/25")),
        ("h", _pk("phoenix Pharmahandel", 5, "WH/IN/31")),
        ("i", _pk("PHOENIX Pharmahandel GmbH & Co KG", 2, "WH/IN/9")),
        ("j", _pk("Gehe-Otto Ltd", 1, "WH/IN/26")),        # 开头是 Gehe → 归 GEHE，不是 Otto
        ("k", _pk("Pharma Otto Ltd", 4, "WH/IN/27")),      # 含 Otto 但不打头 → 其他；与 Queisser 合计同为 4，按名称排前
    ])
    got = [[pid for pid, _ in g] for g in L.group_by_vendor(pickings)]
    assert got == [["h", "i", "d"], ["g"], ["j"], ["b"], ["k"], ["c", "f"], ["a"], ["e"]]


def test_missing_column_is_reported(tmp_path):
    wb = Workbook()
    wb.active.append(["Operations/ID", "Operations/Product"])
    p = tmp_path / "bad.xlsx"
    wb.save(p)
    with pytest.raises(L.InputError, match="move_id"):
        L.run(str(p), str(tmp_path))


def _id_export(path, id_cols):
    """最小导出：一张单一行；`id_cols` = [(列名, 单头值)]，模拟模板里 External ID 被改名成 `ID`。"""
    head = ["Operations/ID", "Operations/Stock Move/ID", "Operations/Product",
            "Operations/Unit of Measure", "Operations/From", "Operations/To",
            "Operations/Stock Move/Demand", "Reference"] + [h for h, _ in id_cols]
    wb = Workbook()
    ws = wb.active
    ws.append(head)
    ws.append([X("ml1"), X("mv1"), "[A_1] A", "Units", "V", "WH/Stock", 2, "WH/IN/1"]
              + [v for _, v in id_cols])
    ws.append([X("ml2"), X("mv2"), "[B_2] B", "Units", "V", "WH/Stock", 1, None] + [None for _ in id_cols])
    wb.save(path)


@pytest.mark.parametrize("id_cols", [
    [("ID", X("stock_picking_433727_9a35dba9")), ("ID", "433727")],   # 2026-10-13 实测：两列同名
    [("ID", "433727"), ("ID", X("stock_picking_433727_9a35dba9"))],   # 顺序反过来也挑对
    [("ID", X("stock_picking_433727_9a35dba9"))],                     # 只有改了名的那列
])
def test_picking_id_column_named_ID(tmp_path, id_cols):
    p = tmp_path / "exp.xlsx"
    _id_export(p, id_cols)
    pk = L.read_export(load_workbook(p).active)
    assert list(pk) == [X("stock_picking_433727_9a35dba9")]
    assert len(pk[X("stock_picking_433727_9a35dba9")]["moves"]) == 2


def test_numeric_id_only_is_rejected(tmp_path):
    p = tmp_path / "exp.xlsx"
    _id_export(p, [("ID", 433727)])
    with pytest.raises(L.InputError, match="database numbers.*picking_id"):
        L.run(str(p), str(tmp_path))
