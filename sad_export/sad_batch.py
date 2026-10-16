"""JD 直邮小包裹荷兰出口申报单（SAD）批次处理：挑问题包裹、整理正常件、填 Export statement。

用法（在仓库根目录跑）：
    python3 -m sad_export.sad_batch verify <上一批 待处理 目录> <上一批 Export statement.xlsx>
    python3 -m sad_export.sad_batch run <源文件夹> <输出目录> [--template <模板.xlsx>]

源文件夹 = 下载解压后的一批：`declaration.xlsx`（TrackingId | DeclarationUcr | FileName）+ 每行一份
`SAD-<TrackingId>-<时间戳>.pdf`（正常单 3 份，补发时 4~6 份）。**源文件夹只读，产出一律复制**。

run 产出（输出目录必须不存在，不覆盖）：
    <输出>/待处理/                            问题包裹 PDF（declaration 里 TrackingId 只出现 1 次）
    <输出>/正常/                              每个 TrackingId 一份带 "Exit Confirmed" 水印的 PDF（多份取时间戳最新）
    <输出>/Export statement - <输出目录名>.xlsx 复制模板，只填「无水印需要POE声明-VP」，Sheet2 不动
模板默认取 <输出目录的上级>/Export statement templete.xlsx。

问题包裹的判定是「只出现 1 次」，但脚本同时解每份 PDF 的内容流查水印，要求两者**完全一致**：
2026-10 批 5274 单里 20 单只出现 1 次且都无水印，其余每单都有水印。识别水印不靠文件名——
那批水印件恒为按时间戳排序的第 3 份、且是同单最大文件，但这只是规律，脚本只拿「最大」作交叉核对提示。

以下任一命中 → 退出码 2、不写任何文件、列出原因（交给人决定，不自行处理）：
    缺 declaration.xlsx / 列名不同 / 与磁盘 PDF 对不上；「只出现 1 次」与「无水印」不一致；
    问题包裹占比 > 3%（2026-09、2026-10 两批分别 1.0%、0.38%）；
    问题包裹的 SAD 不止 1 个 Artikelen（第 5 栏 ≠ 1）、字段取不到、格式异常、两种解析结果不一致。

字段（每个问题包裹一行，按文件名排序）：
    B Parcel No.  第 31 栏「1 - CT」上一行的唛头      C Content  第 31/2 栏品名
    D 金额        页底 `Factuurbedrag: 12,56 EUR`    E 重量     第 35 栏 Brutomassa(kg)
    D、E 照抄德式逗号小数、以文本写入；F MAWB / G Exit Date 留空；H 文件名去 `.pdf`。
主路径是 `pdftotext -layout`（系统 poppler；Mac `brew install poppler`），PDF 内容流 Tm/Tj 坐标是独立
交叉核对。Windows 的 mingw pdftotext 打不开中文路径，统一先复制到临时 ASCII 路径再跑。

verify 用上一批成品回归解析，必须逐格一致；PDF 版式可能变了时先跑它。
"""
import argparse
import copy
import re
import shutil
import subprocess
import sys
import tempfile
import zlib
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.formatting.formatting import ConditionalFormattingList

SHEET = '无水印需要POE声明-VP'
DECL_COLS = ('TrackingId', 'DeclarationUcr', 'FileName')
WATERMARK = b'(Exit Confirmed)Tj'
MAX_RATIO = 0.03
FIELDS = ('parcel', 'content', 'value', 'weight', 'articles')
TOK = re.compile(rb'1 0 0 1 ([\d.]+) ([\d.]+) Tm\n[^\n]*\n\(((?:\\.|[^\\)])*)\)Tj')
ITEM = re.compile(r'^\s+\d+\s+\d+ [A-Z]{2} \S+')  # 第 32 栏商品行「1   1 CT 100016709417 …」


def stream_text(data):
    """解压 PDF 全部内容流。不靠 endstream 截断（压缩数据里可能碰巧含这串字节，
    2026-10-02 因此误判 3 单），从 `>> stream` 之后交给 zlib，由它自己判断流在哪结束。"""
    out = b''
    for m in re.finditer(rb'>>\s*stream\r?\n', data):
        try:
            out += zlib.decompressobj().decompress(data[m.end():])
        except zlib.error:
            pass
    return out


def stream_fields(txt):
    """按内容流里的固定坐标取字段（交叉核对用）"""
    toks = [(float(x), float(y), s.decode('latin1').strip()) for x, y, s in TOK.findall(txt)]

    def at(y, x0=0, x1=9999):
        return next((s for x, yy, s in toks if abs(yy - y) < 1 and x0 <= x < x1 and s), '')

    fact = next((s for _, _, s in toks if s.startswith('Factuurbedrag:')), '')
    return {'parcel': at(430.3333, 60, 70), 'content': at(666.3333, 400, 500),
            'value': fact.replace('Factuurbedrag:', '').replace('EUR', '').strip(),
            'weight': at(337.3333, 550), 'articles': at(119.3333, 420, 480)}


def _seg(line, col, stop=None):
    """一行里起点落在 [col-3, stop) 的那段文字（按 2 个以上空格切段）"""
    for m in re.finditer(r'\S+(?: \S+)*', line):
        if col - 3 <= m.start() < (stop or 10**6):
            return m.group()
    return ''


def _below(lines, i, col, stop=None):
    return next((v for ln in lines[i:i + 6] for v in [_seg(ln, col, stop)] if v), '')


def parse_layout(lines):
    """pdftotext -layout 的文本行 → 字段 + 件数行 / 商品行计数"""
    r = dict.fromkeys(FIELDS, '') | {'colli': 0, 'items': 0}
    for i, ln in enumerate(lines):
        if 'Artikelen (5)' in ln and not r['articles']:
            r['articles'] = _below(lines, i + 1, ln.index('Artikelen (5)'), ln.index('Totaal colli'))
        if 'Brutomassa(kg) (35)' in ln and not r['weight']:
            r['weight'] = _below(lines, i + 1, ln.index('Brutomassa(kg) (35)'))
        if re.search(r'\s\d+ - [A-Z]{2}\s*$', ln):  # 第 31 栏「1 - CT」，上一行是唛头
            r['colli'] += 1
            r['parcel'] = r['parcel'] or lines[i - 1].split()[-1]
        if 'Omschrijving van de goederen (31/2)' in ln and not r['content']:
            col = ln.index('Omschrijving van de goederen (31/2)')
            j = next((k for k in range(i + 1, len(lines)) if ITEM.match(lines[k])), None)
            r['items'] = sum(bool(ITEM.match(x)) for x in lines[i + 1:])
            # 品名多数与商品行同行，有时因纵坐标差一点被排到上一行；收件人名隔一空行，取不到
            if j is not None:
                r['content'] = next((v for k in (j, j - 1, j + 1) for v in [_seg(lines[k], col)] if v), '')
        m = re.search(r'Factuurbedrag:\s*(\S+)\s*EUR', ln)
        if m:
            r['value'] = m.group(1)
    return r


def pdftotext_lines(pdf, tmp):
    t = Path(tmp) / 'x.pdf'
    shutil.copyfile(pdf, t)
    out = subprocess.run(['pdftotext', '-layout', str(t), '-'], capture_output=True, check=True).stdout
    return out.decode('latin1').splitlines()


def check(a, b):
    """a = layout 结果，b = 内容流结果 → 问题列表（空 = 可填表）"""
    issues = [f'{k}: layout={a[k]!r} 内容流={b[k]!r}' for k in FIELDS if a[k] != b[k]]
    issues += [f'{k} 为空' for k in FIELDS[:4] if not a[k]]
    if a['articles'] != '1':
        issues.append(f"Artikelen={a['articles']!r}（多品名，一行填不下）")
    if a['colli'] != 1 or a['items'] != 1:
        issues.append(f"件数行 {a['colli']} / 商品行 {a['items']}")
    if a['value'] and not re.fullmatch(r'\d+,\d{2}', a['value']):
        issues.append(f"金额格式 {a['value']!r}")
    if a['weight'] and not re.fullmatch(r'\d+,\d{3}', a['weight']):
        issues.append(f"重量格式 {a['weight']!r}")
    return issues


def extract(pdf, txt, tmp):
    a = parse_layout(pdftotext_lines(pdf, tmp))
    return {**a, 'issues': check(a, stream_fields(txt))}


def row_values(name, f):
    return ['=ROW()-1', f['parcel'], f['content'], f['value'], f['weight'], None, None, name[:-4]]


def write_statement(template, rows, out):
    """rows: [(文件名, 字段)]。表头 H1 写 File Name；去掉模板 B/D 列的「重复值」高亮
    （唛头本来就会重复，留着会整列标红），C 列 wine/champagne 文字规则保留。"""
    wb = openpyxl.load_workbook(template)
    ws = wb[SHEET]
    ws['H1'] = 'File Name'
    style = [copy.copy(c._style) for c in ws[2][:8]]
    ws.delete_rows(2, ws.max_row)
    cf = ConditionalFormattingList()
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            if rule.type != 'duplicateValues':
                cf.add(str(rng.sqref), rule)
    ws.conditional_formatting = cf
    for i, (name, f) in enumerate(rows, start=2):
        for col, (v, st) in enumerate(zip(row_values(name, f), style), start=1):
            ws.cell(i, col, v)._style = st
        ws.row_dimensions[i].height = 18
    wb.save(out)


def _stop(msg):
    print('停：' + msg, file=sys.stderr)
    sys.exit(2)


def read_declaration(src):
    decl = src / 'declaration.xlsx'
    if not decl.exists():
        _stop(f'{src} 没有 declaration.xlsx')
    rows = list(openpyxl.load_workbook(decl, read_only=True).worksheets[0].iter_rows(values_only=True))
    if tuple(rows[0][:3]) != DECL_COLS:
        _stop(f'declaration.xlsx 列名不同：{rows[0]}（应为 {DECL_COLS}）')
    return [tuple(r[:3]) for r in rows[1:] if r[0]]


def verify(todo, ref):
    ws = openpyxl.load_workbook(ref)[SHEET]
    want = [[c.value for c in row[:8]] for row in ws.iter_rows(min_row=2)]
    got, flagged = [], 0
    with tempfile.TemporaryDirectory() as tmp:
        for p in sorted(Path(todo).glob('*.pdf')):
            f = extract(p, stream_text(p.read_bytes()), tmp)
            flagged += bool(f['issues'])
            for s in f['issues']:
                print('  !', p.name, s)
            got.append(row_values(p.name, f))
    diffs = [(i, 'ABCDEFGH'[j], x, y) for i, (a, b) in enumerate(zip(want, got), 2)
             for j, (x, y) in enumerate(zip(a, b)) if x != y]
    for d in diffs:
        print('  DIFF 行%d %s 参考=%r 提取=%r' % d)
    ok = len(want) == len(got) and not diffs and not flagged
    print(f'参考 {len(want)} 行，提取 {len(got)} 行，差异 {len(diffs)} 格 →', '一致' if ok else '不一致')
    return ok


def run(src, out, template):
    if out.exists():
        _stop(f'输出目录已存在 {out}，不覆盖')
    if not template.exists():
        _stop(f'找不到模板 {template}')
    rows = read_declaration(src)
    stop = []
    on_disk = {p.name for p in src.glob('*.pdf')}
    listed = [r[2] for r in rows]
    if set(listed) != on_disk or len(listed) != len(set(listed)):
        stop.append(f'declaration 与磁盘 PDF 对不上：缺 {sorted(set(listed) - on_disk)[:5]}，'
                    f'多 {sorted(on_disk - set(listed))[:5]}，重复行 {len(listed) - len(set(listed))}')
    stop += [f'文件名与 TrackingId 不符：{r}' for r in rows if not r[2].startswith(f'SAD-{r[0]}-')][:5]

    by_tid = defaultdict(list)
    for tid, _, name in rows:
        by_tid[tid].append(name)
    print(f'declaration {len(rows)} 行，TrackingId {len(by_tid)}，'
          f'份数分布 {sorted(Counter(map(len, by_tid.values())).items())}')

    single = sorted(fs[0] for fs in by_tid.values() if len(fs) == 1)
    wm, fields = {}, {}
    with tempfile.TemporaryDirectory() as tmp:
        for k, name in enumerate(sorted(on_disk & set(listed)), 1):
            txt = stream_text((src / name).read_bytes())
            wm[name] = WATERMARK in txt
            if len(by_tid[name.split('-')[1]]) == 1:
                fields[name] = extract(src / name, txt, tmp)
            if k % 2000 == 0:
                print(f'  已扫 {k}/{len(on_disk)}')

    nowm = {t for t, fs in by_tid.items() if not any(wm.get(f) for f in fs)}
    if {f.split('-')[1] for f in single} != nowm:
        stop.append(f'「只出现 1 次」与「无水印」不一致：只 1 次但有水印 {[f for f in single if wm.get(f)]}；'
                    f'多次但都无水印 {sorted(t for t in nowm if len(by_tid[t]) > 1)}')
    if len(single) > MAX_RATIO * len(by_tid):
        stop.append(f'问题包裹 {len(single)}/{len(by_tid)} 超过 {MAX_RATIO:.0%}，这批的规则可能不同')
    stop += [f'{f}: {s}' for f in single if f in fields for s in fields[f]['issues']]

    size = lambda f: (src / f).stat().st_size
    not_largest = [t for t, fs in by_tid.items() if any(wm.get(f) for f in fs) and not all(wm.get(f) for f in fs)
                   and min(size(f) for f in fs if wm.get(f)) <= max(size(f) for f in fs if not wm.get(f))]
    if not_largest:
        print(f'提示：{len(not_largest)} 单的水印件不是同单最大文件（只作交叉核对，不拦），抽查：{not_largest[:5]}')

    keep = sorted(max((f for f in fs if wm.get(f)), key=lambda f: f.rsplit('-', 1)[1])
                  for fs in by_tid.values() if any(wm.get(f) for f in fs))
    print(f'问题包裹 {len(single)}，正常 {len(keep)}（另有 {sum(wm.values()) - len(keep)} 份重复水印件不收）')
    for f in single:
        r = fields.get(f, {})
        print(f"  {f[:-4]} | {r.get('parcel')} | {r.get('content')} | {r.get('value')} | {r.get('weight')}")
    if stop:
        print('\n停，未写任何文件：')
        for s in stop:
            print('  -', s)
        sys.exit(2)

    for sub, files in (('待处理', single), ('正常', keep)):
        (out / sub).mkdir(parents=True)
        for f in files:
            shutil.copy2(src / f, out / sub / f)
    xlsx = out / f'Export statement - {out.name}.xlsx'
    write_statement(template, [(f, fields[f]) for f in single], xlsx)

    a, b = (len(list((out / s).glob('*.pdf'))) for s in ('待处理', '正常'))
    print(f'\n完成 → {xlsx}\n  待处理 {a} + 正常 {b} = {a + b}，TrackingId {len(by_tid)}，守恒={a + b == len(by_tid)}；'
          f'源文件夹仍 {len(list(src.glob("*.pdf")))} 份（未动）')


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = ap.add_subparsers(dest='cmd', required=True)
    v = sub.add_parser('verify', help='用上一批成品回归 PDF 解析，必须逐格一致')
    v.add_argument('todo', type=Path, help='上一批 待处理 目录')
    v.add_argument('ref', type=Path, help='上一批 Export statement.xlsx')
    r = sub.add_parser('run', help='处理一批')
    r.add_argument('src', type=Path, help='源文件夹（含 declaration.xlsx 与 SAD PDF，只读）')
    r.add_argument('out', type=Path, help='输出目录（不能已存在）')
    r.add_argument('--template', type=Path, help='默认 <输出目录的上级>/Export statement templete.xlsx')
    a = ap.parse_args(argv)
    if a.cmd == 'verify':
        sys.exit(0 if verify(a.todo, a.ref) else 1)
    run(a.src, a.out, a.template or a.out.parent / 'Export statement templete.xlsx')


if __name__ == '__main__':
    main()
