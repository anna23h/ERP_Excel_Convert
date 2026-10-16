"""sad_batch 测试：layout 文本解析、内容流解压 / 水印 / 坐标取值、校验、按模板写表。

不读真 PDF（含收件人信息，不进仓库）：layout 行照 2026-10 实测 pdftotext -layout 的列位置构造，
收件人用假名；内容流用 zlib 现压；模板用代码造（*.xlsx 不进仓库）。端到端回归走 `verify` 子命令。
"""
import os
import sys
import zlib

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule, Rule

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from sad_export import sad_batch as S  # noqa: E402

C1, C2 = 35, 129  # 左栏 / 右栏起始列


def L(left='', right='', rcol=C2):
    """左栏文字从第 35 列起、右栏从 rcol 列起的一行"""
    return (' ' * C1 + left).ljust(rcol) + right if right else ' ' * C1 + left


def page(content_offset=0, articles='1', colli=('1 - CT',), value='12,56', weight='0,098'):
    """content_offset：0 = 品名与商品行同行，-1 = 品名排到商品行上一行（实测两种都有）"""
    item = L('1            1 CT 100016709417')
    lines = [
        ' ' * C2 + 'Artikelen (5) Totaal colli (6)           Datum van afgifte: 31-07-2026 16:24',
        ' ' * 145 + 'Douanekantoor: NL000432',
        ' ' * (C2 + 8) + articles + '                 1',
        L('Identiteit en nationaliteit van het vervoermiddel bij vertrek (18)',
          'Brutomassa(kg) (35)', 170),
        L('---', weight, 190),
        'Colli en om-' + ' ' * 23 + 'Merken en nummers - Container (s) nr (s) - Aantal en soort',
        'schrijving' + ' ' * 25 + '100016709417',
        *['van de' + ' ' * 29 + c for c in colli],
        L('Art. nr. (32) Aantal en soort colli', 'Omschrijving van de goederen (31/2)'),
        L('Afzender/Exporteur (2)', 'Geadresseerde (8)'),
        L('UNDG (44/4)', 'Soort aangifte( 1)    Statistische waarde (46)'),
    ]
    if content_offset == 0:
        lines += [item.ljust(C2) + 'Nasal congestion relief oral drops']
    else:
        lines += [' ' * C2 + 'Nasal congestion relief oral drops', item]
    lines += ['', ' ' * C2 + 'test receiver', ' ' * C2 + '00000', '',
              L('Klantreferentie:', f'Factuurbedrag: {value} EUR', 170)]
    return lines


@pytest.mark.parametrize('offset', [0, -1])
def test_parse_layout(offset):
    r = S.parse_layout(page(offset))
    assert (r['parcel'], r['content'], r['value'], r['weight'], r['articles']) == \
        ('100016709417', 'Nasal congestion relief oral drops', '12,56', '0,098', '1')
    assert (r['colli'], r['items']) == (1, 1)


def ok_stream():
    return {'parcel': '100016709417', 'content': 'Nasal congestion relief oral drops',
            'value': '12,56', 'weight': '0,098', 'articles': '1'}


def test_check_clean():
    assert S.check(S.parse_layout(page()), ok_stream()) == []


@pytest.mark.parametrize('kw, expect', [
    ({'articles': '2'}, 'Artikelen'),
    ({'colli': ('1 - CT', '1 - CT')}, '件数行'),
    ({'value': '12.56'}, '金额格式'),
    ({'weight': '0,1'}, '重量格式'),
])
def test_check_flags(kw, expect):
    a = S.parse_layout(page(**kw))
    b = {k: a[k] for k in S.FIELDS}  # 内容流与 layout 一致，只看格式类问题
    assert any(expect in s for s in S.check(a, b))


def test_check_flags_parser_disagreement():
    b = ok_stream() | {'content': 'something else'}
    assert any(s.startswith('content:') for s in S.check(S.parse_layout(page()), b))


def tj(x, y, s):
    return f'BT\n1 0 0 1 {x} {y} Tm\n/F1 8 Tf\n({s})Tj\nET\n'.encode('latin1')


def fake_pdf(*streams):
    return b'%PDF-1.4\n' + b''.join(
        b'<< /Length 9 /Filter /FlateDecode >> stream\n' + zlib.compress(s) + b'\nendstream\n' for s in streams)


def test_stream_watermark_and_fields():
    body = (tj(65, 430.3333, '100016709417') + tj(450, 666.3333, 'Eye Drops') + tj(560, 337.3333, '0,005')
            + tj(440, 119.3333, '1') + tj(300, 40, 'Factuurbedrag: 10,19 EUR'))
    txt = S.stream_text(fake_pdf(body, b'q (Exit Confirmed)Tj Q'))
    assert S.WATERMARK in txt
    assert S.stream_fields(txt) == {'parcel': '100016709417', 'content': 'Eye Drops', 'value': '10,19',
                                    'weight': '0,005', 'articles': '1'}


def test_stream_survives_endstream_bytes_inside_data():
    # 压缩数据里碰巧出现 endstream 时，按 endstream 截断会解压失败；zlib 自己判断结束则不受影响
    payload = b'q (Exit Confirmed)Tj Q'
    raw = b'<< /Length 9 >> stream\n' + zlib.compress(payload) + b'endstream junk'
    assert S.WATERMARK in S.stream_text(raw)


def test_write_statement(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = S.SHEET
    ws.append(['NO.', 'Parcel No.', 'Content', 'Total Value In Invoice (EUR)', 'Declared Weight (KG)',
               'MAWB Number', 'Exit Date (UTC)', '文件名'])
    for i in range(3):
        ws.append(['=ROW()-1', f'old{i}', 'old', '1,00', 0.08, None, None, f'old{i}'])
    ws['D2'].number_format = '0.00_ '
    ws.conditional_formatting.add('B2:B27', Rule(type='duplicateValues'))
    ws.conditional_formatting.add('C2:C6', FormulaRule(formula=['ISNUMBER(SEARCH("wine",C2))']))
    wb.create_sheet('Sheet2')['A1'] = 'keep me'
    tpl, out = tmp_path / 't.xlsx', tmp_path / 'o.xlsx'
    wb.save(tpl)

    f = S.parse_layout(page())
    S.write_statement(tpl, [('SAD-JD1-2.pdf', f), ('SAD-JD3-4.pdf', f)], out)
    got = load_workbook(out)
    ws = got[S.SHEET]
    assert ws['H1'].value == 'File Name'
    assert [[c.value for c in r] for r in ws.iter_rows(min_row=2)] == [
        ['=ROW()-1', '100016709417', 'Nasal congestion relief oral drops', '12,56', '0,098', None, None, n]
        for n in ('SAD-JD1-2', 'SAD-JD3-4')]
    assert ws['D3'].number_format == '0.00_ ' and ws['D3'].data_type == 's'
    rules = [r.type for rng in ws.conditional_formatting for r in rng.rules]
    assert 'duplicateValues' not in rules and rules  # 去掉重复值高亮，保留其他规则
    assert got['Sheet2']['A1'].value == 'keep me'


def test_cli_requires_subcommand():
    with pytest.raises(SystemExit):
        S.main([])
