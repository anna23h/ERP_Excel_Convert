#!/usr/bin/env python3
"""天猫面单 PDF 按「拣货表+面单」顺序重排、合并成一份（ISSUES: [拣货面单] L）。

天猫批量打印下载的面单 PDF 每份最多 100 页，页序与阶段一产出的「面单」sheet 序号对不上。
本脚本按每页上的 LP 运单号（= 面单 sheet 的 VO Tracking No）重排，合并成一份 PDF。

用法:
    python vo_orders/sort_labels.py <拣货表+面单.xlsx> <面单.pdf> [<面单.pdf> ...] [-o 输出目录]
默认输出到 xlsx 同目录：…{店}{n}单 面单（按拣货表排序）.pdf；原 PDF / xlsx 只读。

两边对不上就全停、不出 PDF（退出码 2）：某页没有 LP 或有多个；同一运单两页；
Excel 有 PDF 缺；PDF 有 Excel 无。多品订单的续行同运单号只算一次，取首次出现的位置。
"""
import argparse
import os
import re
import sys

from openpyxl import load_workbook

FACE_SHEET = "面单"
TRACKING_COL = "VO Tracking No"
LP_RE = re.compile(r"LP\d{14}")
OUT_SUFFIX = "面单（按拣货表排序）"


class MismatchError(ValueError):
    """面单 PDF 与拣货表对不上（CLI 退出码 2）。"""


def _pdfium():
    try:
        import pypdfium2
    except ImportError:
        raise ValueError("缺 pypdfium2：pip install -r requirements.txt") from None
    return pypdfium2


def read_order(xlsx):
    """「面单」sheet 的运单号顺序（去重，取首次出现）。按列名取，不依赖列位置。"""
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    if FACE_SHEET not in wb.sheetnames:
        raise ValueError(f"{os.path.basename(xlsx)} 里没有「{FACE_SHEET}」sheet，不是拣货表+面单？")
    rows = wb[FACE_SHEET].iter_rows(values_only=True)
    header = [str(c).strip() if c is not None else "" for c in next(rows, [])]
    if TRACKING_COL not in header:
        raise ValueError(f"「{FACE_SHEET}」sheet 缺「{TRACKING_COL}」列")
    col = header.index(TRACKING_COL)
    order, seen = [], set()
    for r in rows:
        lp = str(r[col]).strip() if col < len(r) and r[col] is not None else ""
        if lp and lp not in seen:
            seen.add(lp)
            order.append(lp)
    wb.close()
    if not order:
        raise ValueError(f"「{FACE_SHEET}」sheet 没有运单号")
    return order


def scan_pdf(path):
    """每页抽到的 LP 号列表（去重保序）。"""
    pdfium = _pdfium()
    doc = pdfium.PdfDocument(path)
    try:
        pages = []
        for page in doc:
            text = page.get_textpage().get_text_bounded()
            pages.append(list(dict.fromkeys(LP_RE.findall(text))))
        return pages
    finally:
        doc.close()


def plan(order, scans):
    """order: Excel 运单顺序；scans: [(pdf 路径, [每页 LP 列表])]（列表而非 dict：同一文件选两次也要报出来）。
    返回 [(pdf 路径, 页下标)]，按 order 排。对不上抛 MismatchError，列全所有问题。"""
    name = os.path.basename
    problems, loc = [], {}
    for path, pages in scans:
        for i, lps in enumerate(pages):
            where = f"{name(path)} 第 {i + 1} 页"
            if len(lps) != 1:
                problems.append(f"{where}：{'没有 LP 运单号' if not lps else '有多个 LP：' + ', '.join(lps)}")
                continue
            lp = lps[0]
            if lp in loc:
                p0, i0 = loc[lp]
                problems.append(f"{lp} 重复：{name(p0)} 第 {i0 + 1} 页 与 {where}")
            else:
                loc[lp] = (path, i)
    want = set(order)
    missing = [lp for lp in order if lp not in loc]
    extra = [lp for lp in loc if lp not in want]
    if missing:
        problems.append(f"拣货表有、PDF 缺 {len(missing)} 单：" + ", ".join(missing))
    for lp in extra:
        p, i = loc[lp]
        problems.append(f"PDF 有、拣货表无：{lp}（{name(p)} 第 {i + 1} 页）")
    if problems:
        raise MismatchError("面单 PDF 与拣货表对不上，未生成 PDF：\n  " + "\n  ".join(problems))
    return [loc[lp] for lp in order]


def out_path(xlsx, outdir=None):
    stem = os.path.splitext(os.path.basename(xlsx))[0]
    stem = stem.replace("拣货表+面单", OUT_SUFFIX) if "拣货表+面单" in stem else f"{stem} {OUT_SUFFIX}"
    return os.path.join(outdir or os.path.dirname(os.path.abspath(xlsx)), stem + ".pdf")


def run(xlsx, pdfs, outdir=None):
    """返回 (输出路径, 日志行)。"""
    pdfium = _pdfium()
    if not pdfs:
        raise ValueError("没有选面单 PDF")
    order = read_order(xlsx)
    seen = set()
    for p in pdfs:
        k = os.path.normcase(os.path.abspath(p))
        if k in seen:
            raise MismatchError(f"同一份 PDF 选了两次，未生成 PDF：{p}")
        seen.add(k)
    scans = [(p, scan_pdf(p)) for p in pdfs]
    picks = plan(order, scans)
    out = out_path(xlsx, outdir)
    if os.path.normcase(os.path.abspath(out)) in seen:
        raise ValueError(f"输出文件就是输入之一，请别把上次合并出的 PDF 再选进来：{out}")
    os.makedirs(os.path.dirname(out), exist_ok=True)
    docs = {p: pdfium.PdfDocument(p) for p in pdfs}
    new = pdfium.PdfDocument.new()
    try:
        for p, i in picks:
            new.import_pages(docs[p], [i])
        new.save(out)
    finally:
        new.close()
        for d in docs.values():
            d.close()
    lines = [f"面单 PDF {len(pdfs)} 份 / {sum(len(pg) for _, pg in scans)} 页，"
             f"拣货表 {len(order)} 单，一一对应",
             f"按拣货表排序的面单 已生成: {out}  ({len(picks)} 页)"]
    return out, lines


def main(argv=None):
    ap = argparse.ArgumentParser(description="天猫面单 PDF 按拣货表+面单顺序重排、合并成一份")
    ap.add_argument("xlsx", help="阶段一产出的 …拣货表+面单.xlsx")
    ap.add_argument("pdfs", nargs="+", help="天猫下载的面单 PDF（可多份）")
    ap.add_argument("-o", "--outdir", help="输出目录（默认 xlsx 同目录）")
    a = ap.parse_args(argv)
    try:
        _, lines = run(a.xlsx, a.pdfs, a.outdir)
    except MismatchError as e:
        print(e, file=sys.stderr)
        return 2
    except (ValueError, OSError) as e:
        print(f"出错：{e}", file=sys.stderr)
        return 1
    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    sys.exit(main())
