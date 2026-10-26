"""sort_labels 测试：运单顺序读取、对不上就全停的护栏、端到端重排合并。

不读真面单：PDF 用代码手写最小文件（每页一行文字），拣货表用 openpyxl 造（*.xlsx/*.pdf 不进仓库）。
"""
import os
import sys

import pytest
from openpyxl import Workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import sort_labels as S  # noqa: E402

pdfium = pytest.importorskip("pypdfium2")


def lp(n):
    return f"LP{n:014d}"


def make_pdf(path, texts):
    """最小 PDF：每页一行 Helvetica 文字。"""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None,
            "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for t in texts:
        stream = f"BT /F1 10 Tf 20 100 Td ({t}) Tj ET".encode()
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream.decode()}\nendstream")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 288 432] "
                    f"/Resources << /Font << /F1 3 0 R >> >> /Contents {len(objs)} 0 R >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offs = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n{o}\nendobj\n".encode()
    x = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += "".join(f"{o:010d} 00000 n \n" for o in offs).encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{x}\n%%EOF\n".encode()
    with open(path, "wb") as f:
        f.write(out)


def make_xlsx(path, trackings):
    wb = Workbook()
    wb.active.title = "拣货表"
    ws = wb.create_sheet("面单")
    ws.append(["序号", "Order Reference", "VO Tracking No", "Internal Reference"])
    for i, t in enumerate(trackings, 1):
        ws.append([i, f"VO_{i}", t, "SKU"])
    wb.save(path)


def test_read_order_dedups_continuation_rows(tmp_path):
    x = tmp_path / "2026年10月26日VO3单 拣货表+面单.xlsx"
    make_xlsx(x, [lp(3), lp(1), lp(1), None, lp(2)])   # 多品订单续行同运单号
    assert S.read_order(x) == [lp(3), lp(1), lp(2)]


def test_read_order_rejects_wrong_workbook(tmp_path):
    x = tmp_path / "a.xlsx"
    Workbook().save(x)
    with pytest.raises(ValueError, match="面单"):
        S.read_order(x)


def test_plan_orders_by_excel():
    scans = [("a.pdf", [[lp(1)], [lp(2)]]), ("b.pdf", [[lp(3)]])]
    assert S.plan([lp(3), lp(1), lp(2)], scans) == [("b.pdf", 0), ("a.pdf", 0), ("a.pdf", 1)]


@pytest.mark.parametrize("scans,needle", [
    ([("a.pdf", [[lp(1)], []])], "a.pdf 第 2 页：没有 LP"),
    ([("a.pdf", [[lp(1), lp(2)]])], "有多个 LP"),
    ([("a.pdf", [[lp(1)]]), ("b.pdf", [[lp(1)]])], f"{lp(1)} 重复：a.pdf 第 1 页 与 b.pdf 第 1 页"),
    ([("a.pdf", [[lp(1)]])], f"PDF 缺 1 单：{lp(2)}"),
    ([("a.pdf", [[lp(1)], [lp(2)], [lp(9)]])], f"PDF 有、拣货表无：{lp(9)}（a.pdf 第 3 页）"),
])
def test_plan_mismatch_stops(scans, needle):
    with pytest.raises(S.MismatchError, match="未生成 PDF") as e:
        S.plan([lp(1), lp(2)], scans)
    assert needle in str(e.value)


def test_out_path_names():
    assert os.path.basename(S.out_path("d/2026年10月26日VO307单 拣货表+面单（后307）.xlsx")) == \
        "2026年10月26日VO307单 面单（按拣货表排序）（后307）.pdf"
    assert S.out_path("d/x.xlsx", "o") == os.path.join("o", "x 面单（按拣货表排序）.pdf")


def test_run_end_to_end(tmp_path):
    x = tmp_path / "2026年10月26日VO4单 拣货表+面单.xlsx"
    make_xlsx(x, [lp(4), lp(2), lp(2), lp(1), lp(3)])
    a, b = tmp_path / "document (1).pdf", tmp_path / "document (2).pdf"
    make_pdf(a, [f"SF {lp(1)} x", lp(2)])
    make_pdf(b, [lp(3), lp(4)])
    out, lines = S.run(str(x), [str(a), str(b)], str(tmp_path / "out"))
    assert os.path.basename(out) == "2026年10月26日VO4单 面单（按拣货表排序）.pdf"
    doc = pdfium.PdfDocument(out)
    got = [S.LP_RE.findall(p.get_textpage().get_text_bounded()) for p in doc]
    doc.close()
    assert got == [[lp(4)], [lp(2)], [lp(1)], [lp(3)]]
    assert "一一对应" in lines[0]


def test_run_mismatch_writes_nothing(tmp_path):
    x = tmp_path / "x.xlsx"
    make_xlsx(x, [lp(1), lp(2)])
    a = tmp_path / "a.pdf"
    make_pdf(a, [lp(1)])
    with pytest.raises(S.MismatchError):
        S.run(str(x), [str(a)], str(tmp_path / "out"))
    with pytest.raises(S.MismatchError, match="选了两次"):
        S.run(str(x), [str(a), str(a)], str(tmp_path / "out"))
    assert not (tmp_path / "out").exists()


def test_cli_exit_codes(tmp_path, capsys):
    x = tmp_path / "x.xlsx"
    make_xlsx(x, [lp(1)])
    a = tmp_path / "a.pdf"
    make_pdf(a, [lp(2)])
    assert S.main([str(x), str(a)]) == 2
    assert S.main([str(tmp_path / "nope.xlsx"), str(a)]) == 1
    make_pdf(a, [lp(1)])
    assert S.main([str(x), str(a)]) == 0
