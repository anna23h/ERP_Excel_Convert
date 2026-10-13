"""so_from_pdf 测试：PDF 文本解析、PZN 匹配裁决、按模板写表。

不连 Odoo、不读真 PDF：解析吃文本、匹配吃假候选表，模板用代码造（*.xlsx 不进仓库）。
文本行照 2026-10-13 实测 Bestellung 的抽取结果构造（行尾粘着右栏联系方式）。
"""
import os
import sys
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from so_import import so_from_pdf as S  # noqa: E402

PAGE1 = """PZN Artikel/Bezeichnung Inhalt Netto Anz. Nettosumme
17604759 PRODUKT A 30ST 11,784 € 192 2.262,53 € Internet:
00427796 PRODUKT B 20ST 7,504 € 192 1.440,77 € www.example.de
05561513 PRODUKT C 40ST 11,032 € 72 794,30 €
Summe 4.497,60 € Serielle DAFÜ:"""
PAGE2 = "PZN Artikel/Bezeichnung Inhalt Netto Anz. Nettosumme\nMit Annahme der vorliegenden Warenbestellung"

COLS = ["ID", "Order Reference", "Customer", S.COL_PID, S.COL_IREF, S.COL_QTY, S.COL_PRICE]


def test_parse_lines_and_summe():
    lines, summe = S.parse_text([PAGE1, PAGE2])
    assert summe == Decimal("4497.60")
    assert [l["pzn"] for l in lines] == ["17604759", "00427796", "05561513"]
    assert lines[0]["netto"] == Decimal("11.784") and lines[0]["qty"] == 192
    assert lines[0]["nsum"] == Decimal("2262.53") and lines[0]["name"] == "PRODUKT A 30ST"


def test_parse_pads_short_pzn():
    text = "427796 PRODUKT B 20ST 7,504 € 192 1.440,77 €\nSumme 1.440,77 €"
    assert S.parse_text([text])[0][0]["pzn"] == "00427796"


@pytest.mark.parametrize("text", [
    "nichts hier",                                                    # 0 行
    "00427796 PRODUKT B 20ST 7,504 € 192 1.440,77 €",                 # 无 Summe
    "00427796 PRODUKT B 20ST 7,504 € 192 1.440,77 €\nSumme 9.999,99 €",  # 合计对不上 = 漏行
])
def test_parse_refuses_unknown_layout(text):
    with pytest.raises(S.StopError):
        S.parse_text([text])


def _lines():
    return S.parse_text([PAGE1])[0]


def test_match_unique_none_multi():
    lines = _lines()
    cands = {"00427796": [("__export__.p_1", "B", False, True)],
             "05561513": [("__export__.p_2", "C", True, True), ("__export__.p_3", "C 2er", True, True)]}
    problems = S.match(lines, cands, {})
    assert lines[1]["xmlid"] == "__export__.p_1"
    assert [(l["pzn"], len(h)) for l, h in problems] == [("17604759", 0), ("05561513", 2)]


def test_match_pick_resolves_multi():
    lines = _lines()[2:]
    cands = {"05561513": [("__export__.p_2", "C", True, True), ("__export__.p_3", "C 2er", True, True)]}
    assert S.match(lines, cands, {"5561513": "__export__.p_2"}) == []
    assert lines[0]["xmlid"] == "__export__.p_2"


def test_match_pick_must_be_a_hit():
    lines = _lines()[2:]
    cands = {"05561513": [("__export__.p_2", "C", True, True)]}
    with pytest.raises(S.StopError):
        S.match(lines, cands, {"05561513": "__export__.p_9"})


def test_exclude_unknown_pzn_stops():
    with pytest.raises(S.StopError):
        S.apply_exclude(_lines(), ["99999999"])
    kept, dropped = S.apply_exclude(_lines(), ["17604759"])
    assert len(kept) == 2 and dropped[0]["pzn"] == "17604759"


def _template(path, head=("__export__.sale_order_1_x", "S00001", "Kunde GmbH")):
    wb = Workbook()
    ws = wb.active
    ws.append(COLS)
    ws.append(list(head) + [False, None, None, None])
    wb.save(path)


def test_build_follows_template(tmp_path):
    tpl, out = tmp_path / "t.xlsx", tmp_path / "o.xlsx"
    _template(tpl)
    lines = _lines()
    for i, l in enumerate(lines):
        l["xmlid"] = f"__export__.p_{i}"
    wb, cols, head = S.read_template(tpl)
    S.build(wb, cols, head, lines, out)
    rows = list(load_workbook(out).active.iter_rows(values_only=True))
    assert list(rows[0]) == COLS
    assert rows[1][:3] == ("__export__.sale_order_1_x", "S00001", "Kunde GmbH")
    assert rows[2][:3] == (None, None, None) and len(rows) == 4
    assert rows[1][4] is None                       # Internal Reference 留空
    assert rows[1][5] == 192
    assert rows[1][6] == pytest.approx(14.02296)    # 11,784 × 1.19 全精度
    assert lines[2]["price"] == Decimal("13.12808")


def test_template_missing_head_stops(tmp_path):
    tpl = tmp_path / "t.xlsx"
    _template(tpl, head=(None, "S00001", "Kunde GmbH"))
    with pytest.raises(S.StopError):
        S.read_template(tpl)
