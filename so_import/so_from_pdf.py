"""客户 Bestellung PDF → Odoo 销售单导入表（填进 Odoo 导出的 SO 模板）。

用法：
    python3 -m so_import.so_from_pdf <Bestellung.pdf> <sale.order 模板.xlsx>
    python3 -m so_import.so_from_pdf B.pdf T.xlsx --exclude <PZN> <PZN>
    python3 -m so_import.so_from_pdf B.pdf T.xlsx --pick <PZN>=__export__.product_product_<…>
    python3 -m so_import.so_from_pdf B.pdf T.xlsx -o /tmp/x      # 输出目录（默认 output/YYYYMMDD）

**PZN 是唯一连接键，绝不按品名兜底**（2026-10-13 用户定）。PDF 的 PZN 补足 8 位，
与 Odoo `product.product.pzn` 精确比对（含已归档品），取该产品的 `__export__` 外部 ID。
任何一行匹配不到或命中多个产品，**整张表不出**、退出码 2、把这些行列出来——
剔除哪行（`--exclude`）、多个里选哪个（`--pick`）是用户的决定，脚本不猜不跳。

为什么只用 Odoo API、不用产品主数据 xlsx：2026-10-13 实测 0730 的主数据快照
13 个 PZN 里查不到 10 个，API 找回其中 8 个（新建的 `Temp_` 临时品）。快照必然过期。

填值口径：
    Order Lines/Product/ID                 ← 匹配到的外部 ID
    Order Lines/Product/Internal Reference ← 留空（已有 ID，第二把匹配键冗余且易报错）
    Order Lines/Quantity                   ← PDF `Anz.`
    Order Lines/Unit Price                 ← PDF `Netto` × 1.19，全精度不四舍五入
列名列序照抄模板；第一行带模板第 2 行的 ID / Order Reference / Customer，其余行这三列留空。

PDF 版式**只支持一家**（明细行 `PZN 名称 Netto € Anz. Nettosumme €`，`Summe` 收尾）。
解析的护栏：抽出的行 Nettosumme 合计必须等于印的 Summe，否则视为版式不认识、拒绝出表——
换了别家版式会在这里停下，而不是静默漏行。
"""
import argparse
import datetime as dt
import os
import re
import sys
from decimal import Decimal

# Windows 控制台默认 cp936/cp1252，打印 × ✓ 这类符号会 UnicodeEncodeError 直接崩，强制 UTF-8。
for _stream in (sys.stdout, sys.stderr):
    if hasattr(_stream, "reconfigure"):
        _stream.reconfigure(encoding="utf-8")

from openpyxl import load_workbook

from common.xlsx import unique_path

VAT = Decimal("1.19")
HEAD_COLS = ("ID", "Order Reference", "Customer")
COL_PID = "Order Lines/Product/ID"
COL_IREF = "Order Lines/Product/Internal Reference"
COL_QTY = "Order Lines/Quantity"
COL_PRICE = "Order Lines/Unit Price"

#: 明细行：PZN 名称 净价(三位小数) € 数量 行净额 €；行尾常粘着页面右栏的联系方式，不锚行尾。
ROW_RE = re.compile(r"^(\d{6,8}) (.+?) (\d{1,3}(?:\.\d{3})*,\d{2,4}) € (\d+) (\d{1,3}(?:\.\d{3})*,\d{2}) €")
SUM_RE = re.compile(r"^Summe (\d{1,3}(?:\.\d{3})*,\d{2}) €")


class StopError(Exception):
    """要停下来交给用户裁决的情况：不出表，退出码 2。"""


def de_num(s):
    """德式数字 `2.262,53` → Decimal('2262.53')。"""
    return Decimal(s.replace(".", "").replace(",", "."))


def pzn8(v):
    s = str(v).strip()
    return s.zfill(8) if s.isdigit() else s


def parse_text(page_texts):
    """各页文本 → (明细行列表, Summe)。行内含 pzn/name/netto/qty/nsum/page。"""
    lines, summe = [], None
    for i, text in enumerate(page_texts, 1):
        for ln in (text or "").splitlines():
            m = ROW_RE.match(ln)
            if m:
                lines.append(dict(page=i, pzn=pzn8(m[1]), name=m[2], netto=de_num(m[3]),
                                  qty=int(m[4]), nsum=de_num(m[5])))
                continue
            m = SUM_RE.match(ln)
            if m:
                summe = de_num(m[1])
    if not lines:
        raise StopError("PDF 里没解析出任何明细行——版式不认识（本脚本只支持一家的 Bestellung 版式）。")
    if summe is None:
        raise StopError("PDF 里找不到 `Summe` 行——版式不认识，不敢确定明细是否抽全。")
    printed = sum(l["nsum"] for l in lines)
    if printed != summe:
        raise StopError(f"抽出的 {len(lines)} 行 Nettosumme 合计 {printed} ≠ 印的 Summe {summe}，"
                        "说明有行没抽到或抽错了，不出表。")
    return lines, summe


def read_pdf(path):
    try:
        import pdfplumber
    except ImportError:
        sys.exit("缺 pdfplumber：pip install -r requirements-pdf.txt")
    with pdfplumber.open(path) as pdf:
        return [pg.extract_text() for pg in pdf.pages]


def apply_exclude(lines, exclude):
    """剔除用户指定的 PZN；指定了 PDF 里没有的 PZN 视为笔误，停下。"""
    exclude = {pzn8(p) for p in exclude}
    unknown = exclude - {l["pzn"] for l in lines}
    if unknown:
        raise StopError(f"--exclude 里的 PZN 不在 PDF 上：{sorted(unknown)}")
    return ([l for l in lines if l["pzn"] not in exclude],
            [l for l in lines if l["pzn"] in exclude])


def match(lines, candidates, pick):
    """candidates: pzn → [(外部ID, 品名, sale_ok, active)]。pick: pzn → 外部ID。

    命中恰好 1 个才填进行里；否则收进 problems 原样返回，不猜。
    """
    pick = {pzn8(k): v for k, v in pick.items()}
    unknown = set(pick) - {l["pzn"] for l in lines}
    if unknown:
        raise StopError(f"--pick 里的 PZN 不在（剔除后的）PDF 明细上：{sorted(unknown)}")
    problems = []
    for l in lines:
        hits = candidates.get(l["pzn"], [])
        if l["pzn"] in pick:
            chosen = [h for h in hits if h[0] == pick[l["pzn"]]]
            if not chosen:
                raise StopError(f"--pick {l['pzn']}={pick[l['pzn']]}：这个外部 ID 不在该 PZN 的命中里 "
                                f"{[h[0] for h in hits]}")
            hits = chosen
        if len(hits) == 1:
            l["xmlid"], l["erp_name"], l["sale_ok"], l["active"] = hits[0]
        else:
            problems.append((l, hits))
    return problems


def odoo_candidates(pzns):
    """Odoo 只读：pzn → [(__export__ 外部 ID, 品名, sale_ok, active)]，含已归档品。"""
    from odoo_api.odoo_client import Odoo
    od = Odoo.connect(context={"active_test": False}, verbose=False)
    prods = od.execute("product.product", "search_read", [[["pzn", "in", sorted(pzns)]]],
                       {"fields": ["id", "name", "pzn", "sale_ok", "active"]})
    xids = od.execute("ir.model.data", "search_read",
                      [[["model", "=", "product.product"], ["module", "=", "__export__"],
                        ["res_id", "in", [p["id"] for p in prods]]]],
                      {"fields": ["res_id", "name"]}) if prods else []
    by_res = {}
    for x in xids:
        by_res.setdefault(x["res_id"], []).append(f"__export__.{x['name']}")
    out = {}
    for p in prods:
        ids = by_res.get(p["id"], [])
        if len(ids) != 1:
            # 没有或有多个导出 ID：模板要的是外部 ID，这种产品没法安全写进表
            raise StopError(f"产品 {p['id']}（PZN {p['pzn']}，{p['name'].strip()}）的 __export__ 外部 ID 有 "
                            f"{len(ids)} 个，没法确定写哪个。先在 Odoo 里导出一次该产品生成外部 ID。")
        out.setdefault(pzn8(p["pzn"]), []).append((ids[0], (p["name"] or "").strip(), p["sale_ok"], p["active"]))
    return out


def read_template(template):
    """→ (workbook, 列名, 头信息)。缺列或第 2 行头信息不全就停。"""
    wb = load_workbook(template)
    ws = wb.active
    cols = [c.value for c in ws[1]]
    missing = [c for c in HEAD_COLS + (COL_PID, COL_IREF, COL_QTY, COL_PRICE) if c not in cols]
    if missing:
        raise StopError(f"模板缺列 {missing}；表头实际是 {cols}")
    head = {c: ws.cell(2, cols.index(c) + 1).value for c in HEAD_COLS}
    if not all(head.values()):
        raise StopError(f"模板第 2 行的 ID / Order Reference / Customer 不全：{head}")
    return wb, cols, head


def build(wb, cols, head, lines, out_path):
    """按模板列名列序写导入表：第一行带头信息，其余行只填明细列。"""
    ws = wb.active
    ws.delete_rows(2, ws.max_row)
    for k, l in enumerate(lines):
        l["price"] = l["netto"] * VAT
        vals = {COL_PID: l["xmlid"], COL_IREF: None, COL_QTY: l["qty"], COL_PRICE: float(l["price"])}
        if k == 0:
            vals.update(head)
        ws.append([vals.get(c) for c in cols])
    wb.save(out_path)


def report(lines, dropped, summe):
    print("\n核对表（PZN / PDF 名称 / 数量 / 净价 / ×1.19 / ERP 品名 / 可销售）")
    for k, l in enumerate(lines, 1):
        flag = "" if l["active"] else "  ⚠ 已归档"
        print(f"{k:>3} | {l['pzn']} | {l['name']} | {l['qty']} | {l['netto']} | {l['price'].normalize()} | "
              f"{l['erp_name']} | {'✓' if l['sale_ok'] else '✗'}{flag}")
    if dropped:
        print("\n已剔除：")
        for l in dropped:
            print(f"    {l['pzn']} | {l['name']} | {l['qty']} | Nettosumme {l['nsum']}")
    full = sum(l["netto"] * l["qty"] for l in lines)
    dsum = sum(l["nsum"] for l in dropped)
    target = summe - dsum
    print(f"\n合计核对：出表行 三位净价×数量 全精度 {full}；PDF Summe {summe}"
          + (f" − 剔除行 {dsum} = {target}" if dropped else "") + f"；差 {target - full}")
    for l in lines:
        diff = l["nsum"] - l["netto"] * l["qty"]
        if diff:
            print(f"    行级四舍五入 {l['pzn']}：{l['netto'] * l['qty']} 印成 {l['nsum']}（{diff:+}）")
    no_sale = [l for l in lines if not l["sale_ok"]]
    if no_sale:
        # 只提示不拦：2026-10-13 实测按外部 ID 导入 SO，8 个 sale_ok=False 的 Temp_ 品照常导入成功
        print(f"\nℹ {len(no_sale)} 个产品未勾可销售（sale_ok=False），不影响按外部 ID 导入（2026-10-13 实测）；"
              f"在 SO 界面手动选产品时搜不到它们：{', '.join(l['pzn'] for l in no_sale)}")


def run(pdf, template, outdir=None, exclude=(), pick=None, candidates_fn=odoo_candidates):
    wb, cols, head = read_template(template)
    lines, summe = parse_text(read_pdf(pdf))
    print(f"PDF 明细 {len(lines)} 行（页 {sorted({l['page'] for l in lines})}），Summe {summe}")
    lines, dropped = apply_exclude(lines, exclude)
    problems = match(lines, candidates_fn({l["pzn"] for l in lines}), pick or {})
    if problems:
        msg = ["以下行匹配异常，未出表（剔除用 --exclude PZN，多选一用 --pick PZN=外部ID）："]
        for l, hits in problems:
            msg.append(f"  {l['pzn']} | {l['name']} | 数量 {l['qty']} | 命中 {len(hits)} 个")
            msg += [f"      {h[0]}  {h[1]}  sale_ok={h[2]}{'' if h[3] else '  已归档'}" for h in hits]
        raise StopError("\n".join(msg))
    outdir = outdir or os.path.join("output", f"{dt.date.today():%Y%m%d}")
    os.makedirs(outdir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(pdf))[0]
    out = unique_path(os.path.join(outdir, f"{head['Order Reference']}_import_{stem}.xlsx"))
    build(wb, cols, head, lines, out)
    print(f"写出 {out}（{len(lines)} 行）；头信息 {head}")
    report(lines, dropped, summe)
    return out


def main():
    ap = argparse.ArgumentParser(description="客户 Bestellung PDF → Odoo 销售单导入表")
    ap.add_argument("pdf", help="客户 Bestellung PDF")
    ap.add_argument("template", help="Odoo 导出的 sale.order 模板 xlsx（第 2 行带 ID / Order Reference / Customer）")
    ap.add_argument("-o", "--outdir", help="输出目录（默认 output/YYYYMMDD）")
    ap.add_argument("--exclude", nargs="+", default=[], metavar="PZN", help="剔除这些 PZN 的行（用户裁决后用）")
    ap.add_argument("--pick", nargs="+", default=[], metavar="PZN=外部ID",
                    help="一个 PZN 命中多个产品时指定用哪个（用户裁决后用）")
    a = ap.parse_args()
    pick = {}
    for kv in a.pick:
        if "=" not in kv:
            ap.error(f"--pick 要写成 PZN=外部ID：{kv}")
        k, v = kv.split("=", 1)
        pick[k.strip()] = v.strip()
    try:
        run(a.pdf, a.template, a.outdir, a.exclude, pick)
    except StopError as e:
        print(f"\n✗ {e}", file=sys.stderr)
        sys.exit(2)


if __name__ == "__main__":
    main()
